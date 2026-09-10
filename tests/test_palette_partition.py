"""Exact paint partition, bounded packing and failure provenance regressions."""
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import numpy as np
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
import component_fidelity as CF
import native_vectors as NV
from test_source_owned_trace import raster_centers

NODE = os.environ.get('EDITABLE_PPT_NODE') or shutil.which('node')


@unittest.skipUnless(NODE, 'existing Node required for real tracing integration')
class PalettePartitionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def inputs(self, pixels, name='fixture'):
        source = self.root/(name+'.png')
        Image.fromarray(pixels.astype('uint8'), 'RGBA').save(source)
        plan = dict(part_id=name, observation='Synthetic isolated visible component; fixed test budget',
            source_sha256=CF.digest(source), source_bbox=[0,0,pixels.shape[1],pixels.shape[0]],
            max_native_paths=100, max_native_commands=10_000)
        return source, plan, self.root/(name+'.svg')

    def run_trace(self, source, plan, output, **kwargs):
        return CF.trace_component(source, output, plan['source_bbox'], node=NODE,
            representation='palette_edges', part_plan=plan, **kwargs)

    def assert_cells(self, output, pixels):
        specs, _, _ = NV.read_vectors(output)
        reconstructed = np.zeros_like(pixels)
        occupancy = np.zeros(pixels.shape[:2], dtype=int)
        for spec in specs:
            mask = raster_centers([spec], pixels.shape[:2])
            color = [int(spec['fill'][i:i+2],16) for i in (0,2,4)]
            reconstructed[mask] = color+[round(spec['alpha'][0]*255)]
            occupancy += mask
        visible = pixels[:,:,3] != 0
        np.testing.assert_array_equal(occupancy, visible.astype(int))
        np.testing.assert_array_equal(reconstructed[visible], pixels[visible])
        return specs

    def test_nested_holes_touching_colors_thin_bridges_and_alpha(self):
        pixels = np.zeros((30,30,4),dtype='uint8')
        pixels[1:29,1:29] = [60,90,120,255]
        pixels[3:27,3:27] = [160,70,90,128]
        pixels[7:23,7:23] = 0
        pixels[12:18,12:18] = [60,90,120,255]
        pixels[10,3:27] = [60,90,120,255]
        source,plan,output = self.inputs(pixels)
        result = self.run_trace(source,plan,output)
        self.assert_cells(output,pixels)
        self.assertFalse(result['partition']['geometry_simplified'])
        self.assertEqual(sum(l['pixels'] for l in result['partition']['layers']),900)
        self.assertEqual(result['quantization']['changed_pixels'],0)
        self.assertEqual(result['practical_editability'],'NOT_VERIFIED')
        self.assertFalse(result['final_artifact_verified'])

    def test_random_partition_has_no_overlap_no_gap_and_no_color_loss(self):
        rng=np.random.default_rng(4711)
        palette=np.array([[0,0,0,0],[30,60,90,255],[190,90,40,255],[80,160,120,100]],dtype='uint8')
        for i in range(6):
            pixels=palette[rng.integers(0,4,(15,18))]
            source,plan,output=self.inputs(pixels,str(i))
            self.run_trace(source,plan,output)
            self.assert_cells(output,pixels)

    def test_full_uniform_fill_and_all_transparent(self):
        pixels=np.full((5,6,4),255,dtype='uint8')
        source,plan,output=self.inputs(pixels)
        self.assertEqual(self.run_trace(source,plan,output)['native_paths'],1)
        self.assert_cells(output,pixels)
        pixels[:,:,3]=0
        source,plan,output=self.inputs(pixels,'transparent')
        with self.assertRaises(ValueError): self.run_trace(source,plan,output)
        self.assertFalse(output.exists())

    def test_stack_preserves_visible_colors_and_real_holes_on_random_sources(self):
        rng=np.random.default_rng(2841)
        palette=np.array([[0,0,0,0],[190,90,40,255],[30,60,90,255],[80,160,120,255]],dtype='uint8')
        for i in range(8):
            pixels=palette[rng.integers(0,4,(12,14))]
            source,plan,output=self.inputs(pixels,'stack'+str(i))
            result=CF.trace_component(source,output,plan['source_bbox'],node=NODE,
                representation='palette_stack',part_plan=plan)
            specs,_,_=NV.read_vectors(output)
            visible=np.zeros_like(pixels);support=np.zeros(pixels.shape[:2],bool)
            for spec in specs:
                mask=raster_centers([spec],pixels.shape[:2])
                visible[mask]=[int(spec['fill'][k:k+2],16) for k in (0,2,4)]+[255]
                support|=mask
            np.testing.assert_array_equal(visible,pixels)
            np.testing.assert_array_equal(support,pixels[:,:,3]>0)
            self.assertEqual(result['partition']['composition'],'opaque_source_color_tree')
            self.assertFalse(result['partition']['semantic_depth_inferred'])
            if i==0:
                repeat=CF.trace_component(source,self.root/'stack-repeat.svg',plan['source_bbox'],node=NODE,
                    representation='palette_stack',part_plan=plan)
                self.assertEqual(result['svg_sha256'],repeat['svg_sha256'])

    def test_stack_rejects_translucency_before_quantization(self):
        pixels=np.full((8,8,4),255,dtype='uint8');pixels[4,4]=[40,50,60,128]
        source,plan,output=self.inputs(pixels)
        with self.assertRaisesRegex(ValueError,'binary alpha'):
            CF.trace_component(source,output,plan['source_bbox'],colors=2,node=NODE,
                representation='palette_stack',part_plan=plan)
        self.assertFalse(output.exists())

    def test_stack_requires_plan_and_honors_complete_edit_budget(self):
        pixels=np.full((8,8,4),255,dtype='uint8');pixels[::2,::2]=[40,50,60,255]
        source,plan,output=self.inputs(pixels)
        with self.assertRaisesRegex(ValueError,'part plan'):
            CF.trace_component(source,output,plan['source_bbox'],node=NODE,representation='palette_stack')
        plan['max_native_commands']=5
        with self.assertRaisesRegex(ValueError,'edit budget'):
            CF.trace_component(source,output,plan['source_bbox'],node=NODE,representation='palette_stack',part_plan=plan)
        self.assertFalse(output.exists())

    def test_budget_failure_keeps_all_commands_and_emits_no_svg(self):
        pixels=np.zeros((20,20,4),dtype='uint8');pixels[::2,::2]=[30,90,180,255]
        source,plan,output=self.inputs(pixels)
        plan['max_native_commands']=20
        with self.assertRaisesRegex(ValueError,'edit budget'):self.run_trace(source,plan,output)
        failure=json.loads(output.with_suffix('.trace.failed.json').read_text())
        self.assertEqual(failure['complexity']['native_profile'],'PASS')
        self.assertEqual(failure['complexity']['budget_status'],'FAIL')
        self.assertGreater(failure['complexity']['native_commands'],20)
        self.assertFalse(output.exists())
        self.assertFalse(output.with_suffix('.trace.json').exists())
        with self.assertRaisesRegex(ValueError,'new'):self.run_trace(source,plan,output)

    def test_part_source_binding_types_and_no_plan_rejected(self):
        pixels=np.full((4,4,4),255,dtype='uint8');source,plan,output=self.inputs(pixels)
        for key,value in [('part_id',''),('source_sha256','0'*64),
                          ('max_native_paths',True),('max_native_commands',0),('unknown',1)]:
            with self.subTest(key=key),self.assertRaises(ValueError):
                self.run_trace(source,dict(plan,**{key:value}),output)
        with self.assertRaises(ValueError):
            CF.trace_component(source,output,[0,0,4,4],node=NODE,representation='palette_edges')
        with self.assertRaises(ValueError):
            CF.trace_component(source,output,[0,0,3,4],node=NODE,representation='palette_edges',part_plan=plan)
        self.assertFalse(output.exists())

    def test_intact_hole_parent_not_split_when_path_limit_exceeded(self):
        pixels=np.zeros((12,12,4),dtype='uint8');pixels[1:11,1:11]=[1,2,3,255];pixels[3:9,3:9]=0
        source,plan,output=self.inputs(pixels)
        with patch.object(NV,'MAX_PATH_CHARS',20),self.assertRaisesRegex(ValueError,'intact'):
            self.run_trace(source,plan,output)
        self.assertFalse(output.exists())

    def test_many_islands_keep_region_mapping_and_native_path_ranges(self):
        pixels=np.zeros((20,20,4),dtype='uint8');pixels[::2,::2]=[30,90,180,255]
        source,plan,output=self.inputs(pixels)
        with patch.object(NV,'MAX_PATH_CHARS',150):result=self.run_trace(source,plan,output)
        layer=next(l for l in result['partition']['layers'] if l['rgba'][3])
        ranges=[p['region_range'] for p in layer['paths']]
        self.assertEqual(ranges[0][0],0);self.assertEqual(ranges[-1][1],100)
        self.assertTrue(all(a[1]==b[0] for a,b in zip(ranges,ranges[1:])))
        self.assertEqual(sum(b-a for a,b in ranges),100)
        self.assert_cells(output,pixels)

    def test_quantization_record_determinism_and_native_roundtrip(self):
        pixels=np.random.default_rng(441).integers(0,256,(12,12,4),dtype='uint8');pixels[:,:,3]=255
        source,plan,output=self.inputs(pixels)
        a=self.run_trace(source,plan,output,colors=8)
        b=self.run_trace(source,plan,self.root/'second.svg',colors=8)
        self.assertEqual(a['svg_sha256'],b['svg_sha256'])
        self.assertEqual(a['quantization'],b['quantization'])
        self.assertGreater(a['quantization']['changed_pixels'],0)
        self.assertFalse(a['quantization']['fidelity_verified'])
        quant=Image.fromarray(pixels).quantize(8,method=Image.Quantize.FASTOCTREE,dither=Image.Dither.NONE).convert('RGBA')
        self.assert_cells(output,np.asarray(quant))
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        NV.add_svg_component(slide,output,x=1,y=1,width=2,height=2,name='test-part')
        deck=self.root/'native.pptx';prs.save(deck)
        actual=Presentation(deck)
        self.assertEqual(len(actual.slides[0].shapes[0].shapes),a['native_paths'])
        with zipfile.ZipFile(deck) as z:
            self.assertFalse(any(p.startswith('ppt/media/') for p in z.namelist()))

    def test_native_import_limits_remain_and_failed_diagnostics_apply_to_smooth(self):
        pixels=np.zeros((12,12,4),dtype='uint8');pixels[2:10,2:10]=[40,80,120,255]
        source,plan,output=self.inputs(pixels)
        with patch.object(NV,'MAX_SVG_BYTES',30),self.assertRaises(NV.SVGProfileError):
            CF.trace_component(source,output,plan['source_bbox'],node=NODE)
        report=json.loads(output.with_suffix('.trace.failed.json').read_text())
        self.assertEqual(report['representation'],'smooth')
        self.assertEqual(report['complexity']['native_profile'],'FAIL')
        self.assertGreater(report['complexity']['svg_bytes'],30)


if __name__=='__main__':unittest.main()
