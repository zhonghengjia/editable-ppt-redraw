"""Source ownership and representation invariants, not semantic/Office sign-off."""
import hashlib
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import numpy as np
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import component_fidelity as CF
import component_geometry as G
import native_vectors as NV

NODE = os.environ.get('EDITABLE_PPT_NODE') or shutil.which('node')
SELECTOR = dict(rgb=[[34,34],[119,119],[170,170]],differences=[],observation='Synthetic exact blue support')


def source_image(mask):
    pixels = np.full((*mask.shape,3),255,dtype=np.uint8)
    pixels[mask] = [34,119,170]
    return Image.fromarray(pixels)


def recipe(seeds):
    return dict(part_id='small-part',source_sha256='1'*64,source_bbox=[0,0,12,12],
        selector_sha256=CF.recipe_digest(SELECTOR), seeds=seeds,observation='Source witnesses in the requested small part')


def raster_centers(specs, shape):
    """Independent nonzero winding evaluation at source pixel centers."""
    yy,xx = np.indices(shape); px,py = xx+.5,yy+.5
    filled = np.zeros(shape,dtype=bool)
    for spec in specs:
        winding = np.zeros(shape,dtype=int)
        start = previous = None
        for command in spec['commands']:
            if command.cmd == 'M':
                start = previous = command.args
                continue
            if command.cmd not in ('L','Z'):
                raise AssertionError('source edges must contain only M/L/Z')
            current = start if command.cmd == 'Z' else command.args
            ax,ay = previous; bx,by = current
            cross = (bx-ax)*(py-ay) - (by-ay)*(px-ax)
            winding += ((ay <= py) & (by > py) & (cross > 0))
            winding -= ((ay > py) & (by <= py) & (cross < 0))
            previous = current
        filled |= winding != 0
    return filled


class OwnershipTests(unittest.TestCase):
    def setUp(self):
        self.mask = np.zeros((12,12),dtype=bool)
        self.mask[1:3,1:3] = True
        self.mask[5:11,5:11] = True
        self.value = recipe([[1,1],[2,2]])

    def own(self,value=None):
        return G.owned_support(self.mask,self.value if value is None else value,'1'*64,[0,0,12,12],CF.recipe_digest(SELECTOR))

    def test_witness_not_largest_component_and_unassigned_ledger(self):
        owned,evidence = self.own()
        self.assertEqual(int(owned.sum()),4)
        self.assertEqual(evidence['unassigned_components'][0]['area'],36)
        self.assertFalse(evidence['semantics_verified'])
        self.assertEqual(int(self.mask.sum()),40, 'source stays immutable')

    def test_source_crop_selector_must_match(self):
        for key,value in [('source_sha256','0'*64),('selector_sha256','0'*64),('source_bbox',[0,0,11,12]),('source_bbox',[False,0,12,12])]:
            with self.subTest(key=key), self.assertRaises(ValueError): self.own(dict(self.value,**{key:value}))

    def test_invalid_background_duplicate_and_disagreeing_witnesses(self):
        for seeds in ([],[[0,0]],[[12,1]],[[1,1.0]],[[True,1]],[[1,1],[1,1]],[[1,1],[6,6]],[[float('nan'),1]],[[1]]):
            with self.subTest(seeds=seeds), self.assertRaises(ValueError): self.own(dict(self.value,seeds=seeds))

    def test_unknown_fields_and_empty_identity_fail(self):
        for value in (None,{},dict(self.value,extra=1),dict(self.value,part_id=''),dict(self.value,observation='')):
            with self.assertRaises(ValueError):
                G.owned_support(self.mask,value,'1'*64,[0,0,12,12],CF.recipe_digest(SELECTOR))

    def test_holes_diagonal_contacts_and_border_membership(self):
        mask = np.ones((7,7),dtype=bool); mask[2:5,2:5]=False
        labels,regions = G.label_regions(mask)
        self.assertTrue(regions[0]['touches_border'])
        self.assertEqual(G.topology(labels==1)['holes'],1)
        self.assertEqual(len(G.label_regions(np.eye(4,dtype=bool))[1]),1)
        self.assertEqual(len(G.label_regions(np.eye(4,dtype=bool),False)[1]),4)

    def test_ownership_limit_refuses_instead_of_dropping_fragments(self):
        mask=np.zeros((90,90),dtype=bool);mask[::2,::2]=True
        value=dict(self.value,source_bbox=[0,0,90,90],seeds=[[0,0]])
        with self.assertRaises(G.GeometryLimit):
            G.owned_support(mask,value,'1'*64,[0,0,90,90],CF.recipe_digest(SELECTOR))


@unittest.skipUnless(NODE,'existing Node required for real tracing integration')
class RepresentationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)

    def tearDown(self): self.temp.cleanup()

    def trace(self,mask,name='candidate',**kwargs):
        source=self.root/f'{name}.png'; source_image(mask).save(source)
        result=CF.trace_component(source,self.root/f'{name}.svg',[0,0,mask.shape[1],mask.shape[0]],
                                  node=NODE,selector=SELECTOR,**kwargs)
        specs,_,_=NV.read_vectors(self.root/f'{name}.svg')
        return result,specs

    def test_exact_cells_for_nested_holes_narrow_bridges_and_border(self):
        nested=np.zeros((25,25),dtype=bool);nested[1:24,1:24]=True;nested[3:22,3:22]=False
        nested[7:18,7:18]=True;nested[10:15,10:15]=False
        bridge=np.zeros((12,12),dtype=bool);bridge[1:5,1:4]=True;bridge[7:11,8:11]=True
        bridge[4,3:9]=True;bridge[4:8,8]=True
        border=np.zeros((12,12),dtype=bool);border[:,0]=True;border[-1,:]=True
        for i,mask in enumerate((nested,bridge,border,np.eye(8,dtype=bool),np.fliplr(np.eye(8,dtype=bool)))):
            with self.subTest(case=i):
                result,specs=self.trace(mask,str(i))
                np.testing.assert_array_equal(raster_centers(specs,mask.shape),mask)
                self.assertEqual(result['representation'],'source_edges')
                self.assertFalse(result['final_artifact_verified'])

    def test_deterministic_random_masks_exact_no_fragment_filter(self):
        rng=np.random.default_rng(281)
        for i in range(24):
            mask=rng.random((9,13))<(.1+.03*i)
            result,specs=self.trace(mask,str(i))
            np.testing.assert_array_equal(raster_centers(specs,mask.shape),mask)
            self.assertEqual(result['selection']['mask_sha256'],hashlib.sha256(mask.tobytes()).hexdigest())

    def test_ownership_integrated_into_preparation_with_hashes(self):
        mask=np.zeros((12,12),dtype=bool);mask[1:3,1:3]=True;mask[5:11,5:11]=True
        source=self.root/'owned.png';source_image(mask).save(source)
        owner=dict(recipe([[1,1]]),source_sha256=CF.digest(source))
        result,specs=self.trace(mask,'owned',ownership=owner)
        self.assertEqual(int(raster_centers(specs,mask.shape).sum()),4)
        selection=result['selection']
        self.assertNotEqual(selection['input_mask_sha256'],selection['mask_sha256'])
        self.assertEqual(selection['ownership']['recipe_sha256'],CF.recipe_digest(owner))
        self.assertEqual(selection['ownership']['unassigned_components'][0]['area'],36)

    def test_explicit_smooth_and_invalid_representation(self):
        mask=np.zeros((12,12),dtype=bool);mask[2:9,2:9]=True
        result,_=self.trace(mask,'smooth',representation='smooth')
        self.assertEqual(result['representation'],'smooth')
        with self.assertRaises(ValueError):self.trace(mask,'bad',representation='magic')
        with self.assertRaises(ValueError):
            CF.trace_component(self.root/'smooth.png',self.root/'bad2.svg',[0,0,12,12],node=NODE,representation='source_edges')
        with self.assertRaises(ValueError):
            CF.trace_component(self.root/'smooth.png',self.root/'bad3.svg',[0,0,12,12],node=NODE,ownership=recipe([[3,3]]))

    def test_native_reopened_contains_custom_paths_not_pictures(self):
        mask=np.zeros((12,12),dtype=bool);mask[1:11,1:11]=True;mask[4:8,4:8]=False
        self.trace(mask)
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        NV.add_svg_component(slide,self.root/'candidate.svg',0,0,2,2,name='owned')
        prs.save(self.root/'native.pptx')
        with zipfile.ZipFile(self.root/'native.pptx') as z:
            xml=z.read('ppt/slides/slide1.xml').decode()
            self.assertIn('custGeom',xml);self.assertIn('lnTo',xml);self.assertNotIn('<p:pic>',xml)
        self.assertEqual(len(Presentation(self.root/'native.pptx').slides),1)

    def test_native_rejection_is_not_silent_fallback(self):
        mask=np.zeros((12,12),dtype=bool);mask[2:9,2:9]=True
        with patch.object(NV,'read_vectors',side_effect=NV.SVGProfileError('test complexity refusal')):
            with self.assertRaises(NV.SVGProfileError):self.trace(mask)
        self.assertFalse((self.root/'candidate.svg').exists())
        self.assertFalse((self.root/'candidate.trace.json').exists())

    def test_worker_rejects_nonbinary_nonuniform_and_empty_support(self):
        for kind in ('partial','colors','empty','full'):
            im=Image.new('RGBA',(4,4),(0,0,0,0))
            if kind=='partial':im.putpixel((1,1),(34,119,170,128))
            if kind=='colors':
                im.putpixel((1,1),(34,119,170,255));im.putpixel((2,2),(255,0,0,255))
            if kind=='full':im=Image.new('RGBA',(4,4),(34,119,170,255))
            with self.assertRaises(ValueError):
                CF.worker(dict(action='trace',representation='source_edges',width=4,height=4,first=CF.rgba(im)),NODE)


if __name__=='__main__': unittest.main()
