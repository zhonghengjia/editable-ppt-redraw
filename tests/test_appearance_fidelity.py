import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import appearance_fidelity as appearance
import component_assets as assets


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Appearance(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.source = self.root/'source.png'
        self.render = self.root/'render.png'
        self.artifact = self.root/'deck.pptx'
        self.image = Image.new('RGB', (100, 100), 'white')
        draw = ImageDraw.Draw(self.image)
        draw.rectangle((10, 10, 29, 29), fill=(60, 60, 60))
        draw.rectangle((60, 10, 79, 29), fill=(200, 200, 200))
        self.image.save(self.source)
        self.image.save(self.render)
        self.make_deck()
        self.probe = dict(kind='luma_delta', samples=[[10,10,20,20], [60,10,20,20]], tolerance=.05, min_pixels=25)
        self.req = dict(id='contrast', dimension='tone', observation='left darker than right in source',
                        certainty='observed', criterion='retain signed local contrast', measurement=self.probe)
        self.item = dict(id='object', source_bbox=[0,0,100,100], comparison='source_exact', requirements=[self.req])
        self.contract = dict(source_sha256=sha(self.source), source_size=[100,100], color_space='srgb', matte=[255,255,255], items=[self.item])
        self.manifest = dict(mode='faithful', source={'path': 'source.png'}, source_inventory=[{'id':'object', 'appearance_sensitive':True}], appearance_fidelity=self.contract)
        self.evidence = dict(slide=1, renderer='independent synthetic test', rendered_from_final_artifact=True,
                             artifact_sha256=sha(self.artifact), render_path='render.png', render_sha256=sha(self.render))

    def make_deck(self, reversed_order=False, duplicate=False, mixed=False, grouped=False):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        for name in (['front', 'back'] if reversed_order else ['back', 'front']):
            if mixed and name == 'back':
                shape = shapes.add_picture(str(self.source), Inches(1), Inches(1), Inches(1), Inches(1))
            else:
                shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
            shape.name = name
        if duplicate:
            shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0,100,100).name='front'
        prs.save(self.artifact)

    def write_evidence(self, review=True):
        self.evidence.update(artifact_sha256=sha(self.artifact), render_sha256=sha(self.render))
        if review:
            self.evidence['appearance_review'] = dict(contract_sha256=appearance.contract_hash(self.contract), reviewer='synthetic assertion fixture', records=[
                dict(item_id=i['id'], requirement_id=r['id'], status='PASS', observation='synthetic fixture comparison',
                     views=[dict(path='render.png', sha256=sha(self.render), scale=s) for s in ('delivered','detail')])
                for i in self.contract['items'] for r in i['requirements']])
        path = self.root/'evidence.json'
        path.write_text(json.dumps(self.evidence), encoding='utf-8')
        return path

    def audit(self, review=True):
        return appearance.audit(self.artifact, self.manifest, self.root/'manifest.json', self.write_evidence(review))

    def occlusion(self, scope='native'):
        self.req.clear()
        self.req.update(id='overlap', dimension='occlusion', observation='front occludes back in source', certainty='observed', criterion='retain local overlap', scope=scope)
        if scope != 'raster_internal':
            self.req['paint_order'] = dict(front_name='front', back_name='back', slide=1)

    def request(self):
        return dict(provider='builtin_imagegen', authorization='synthetic test approval', attempt=1,
                    prompt='One isolated illustration.', subject='synthetic object', style='source faithful',
                    invariants=['source topology'], references=[])

    def invariant(self):
        self.req.pop('measurement', None)
        self.item.update(comparison='invariants', authorization='user permits illustrative replacement')

    def test_source_contrast_control(self):
        result = self.audit()
        self.assertTrue(result['valid']); self.assertFalse(result['unverified'])
        self.assertEqual(result['semantic_review'], 'NOT_VERIFIED')

    def test_flattened_multitone_rejected(self):
        Image.new('RGB',(100,100),(130,130,130)).save(self.render)
        self.assertFalse(self.audit()['valid'])

    def test_reversed_light_dark_rejected(self):
        draw = ImageDraw.Draw(self.image)
        draw.rectangle((10,10,29,29),fill=(200,200,200)); draw.rectangle((60,10,79,29),fill=(60,60,60))
        self.image.save(self.render)
        self.assertFalse(self.audit()['valid'])

    def test_front_brighter_than_back_is_valid(self):
        self.probe['samples'].reverse()
        result = self.audit()
        self.assertTrue(result['valid'])
        self.assertGreater(result['requirements'][0]['measurement']['source'], 0)

    def test_flat_source_control_and_invented_shading(self):
        self.req.update(dimension='flat', measurement=dict(kind='luma_spread', samples=[[0,0,100,100]], tolerance=.05,min_pixels=25))
        Image.new('RGB',(100,100),'white').save(self.source)
        self.contract['source_sha256']=sha(self.source)
        Image.new('RGB',(100,100),'white').save(self.render)
        self.assertTrue(self.audit()['valid'])
        self.image.paste((60,60,60), (0,0,50,100)); self.image.save(self.render)
        self.assertFalse(self.audit()['valid'])

    def test_semantic_color_swap_rejected(self):
        self.req.update(dimension='semantic_color', measurement=dict(kind='rgb_median',samples=[[10,10,20,20]], tolerance=.05,min_pixels=25))
        self.image.paste((0,180,0), (10,10,30,30)); self.image.save(self.render)
        self.assertFalse(self.audit()['valid'])

    def test_missing_review_is_not_pass(self):
        result=self.audit(False)
        self.assertTrue(result['valid']); self.assertTrue(result['unverified'])

    def test_small_samples_unverified(self):
        self.probe['min_pixels']=401
        self.assertTrue(self.audit()['unverified'])

    def test_uncertain_source_cannot_be_certified(self):
        self.req['certainty']='uncertain'
        self.assertTrue(self.audit()['unverified'])

    def test_stale_source_hash_rejected(self):
        self.contract['source_sha256']='a'*64
        self.assertFalse(self.audit()['valid'])

    def test_stale_artifact_and_render_hashes_rejected(self):
        path=self.write_evidence()
        for field in ('artifact_sha256','render_sha256'):
            wrong=copy.deepcopy(self.evidence); wrong[field]='a'*64
            path.write_text(json.dumps(wrong),encoding='utf-8')
            self.assertFalse(appearance.audit(self.artifact,self.manifest,self.root/'manifest.json',path)['valid'])

    def test_stale_review_and_view_hashes_rejected(self):
        path=self.write_evidence()
        for field in ('contract','view'):
            wrong=copy.deepcopy(self.evidence)
            if field=='contract': wrong['appearance_review']['contract_sha256']='a'*64
            else: wrong['appearance_review']['records'][0]['views'][0]['sha256']='a'*64
            path.write_text(json.dumps(wrong),encoding='utf-8')
            self.assertFalse(appearance.audit(self.artifact,self.manifest,self.root/'manifest.json',path)['valid'])

    def test_registration_cannot_hide_changed_canvas(self):
        Image.new('RGB',(200,120),'white').save(self.render)
        self.assertFalse(self.audit()['valid'])

    def test_native_and_grouped_order(self):
        self.occlusion()
        for grouped in (False,True):
            self.make_deck(grouped=grouped)
            self.assertEqual(self.audit()['requirements'][0]['paint_order'],'PASS')
            self.make_deck(reversed_order=True,grouped=grouped)
            self.assertFalse(self.audit()['valid'])

    def test_duplicate_object_names_unverified(self):
        self.occlusion(); self.make_deck(duplicate=True)
        self.assertTrue(self.audit()['unverified'])

    def test_mixed_order_and_wrong_scope(self):
        self.occlusion('mixed'); self.make_deck(mixed=True)
        self.assertFalse(self.audit()['unverified'])
        self.req['scope']='native'
        self.assertTrue(self.audit()['unverified'])

    def test_raster_internals_never_package_order(self):
        self.occlusion('raster_internal')
        self.assertEqual(self.audit()['requirements'][0]['paint_order'],'NOT_APPLICABLE')
        self.req['paint_order']=dict(front_name='front',back_name='back',slide=1)
        self.assertTrue(appearance.validate_contract(self.manifest))

    def test_transparency_requires_background_review(self):
        self.req.pop('measurement'); self.req['dimension']='transparency'
        self.assertTrue(self.audit()['unverified'])

    def test_numeric_probe_not_surrogate_or_redesign(self):
        for mode in ('semantic','redesign'):
            self.manifest['mode']=mode
            self.assertTrue(appearance.validate_contract(self.manifest))
        self.manifest['mode']='faithful'
        self.item.update(comparison='invariants', authorization='approved')
        self.assertTrue(appearance.validate_contract(self.manifest))

    def test_schema_type_bounds_coverage_and_drift(self):
        for field,value in [('source_size',[True,100]),('matte',[0,0,256]),('items',[]),('unknown',1)]:
            m=copy.deepcopy(self.manifest); m['appearance_fidelity'][field]=value
            self.assertTrue(appearance.validate_contract(m))
        self.probe['samples'][0]=[99,0,20,20]
        self.assertTrue(appearance.validate_contract(self.manifest))

    def test_legacy_and_new_coverage_boundary(self):
        self.manifest.pop('appearance_fidelity')
        self.assertTrue(appearance.validate_contract(self.manifest))
        self.manifest['source_inventory'][0].pop('appearance_sensitive')
        self.assertEqual(appearance.validate_contract(self.manifest), [])
        self.assertTrue(appearance.audit(self.artifact,self.manifest,self.root/'manifest.json',None)['unverified'])

    def test_generation_binding_and_contract_drift(self):
        self.invariant()
        request=appearance.bind_generation(self.request(),self.manifest,['object'])
        self.assertTrue(assets.generation_preflight(request,manifest=self.manifest)['valid'])
        self.assertFalse(assets.generation_preflight(request)['valid'])
        self.req['criterion']='retain updated source observation'
        self.assertFalse(assets.generation_preflight(request,manifest=self.manifest)['valid'])

    def test_prompt_drift_and_no_binding_accumulation(self):
        self.invariant()
        request=appearance.bind_generation(self.request(),self.manifest,['object'])
        with self.assertRaises(ValueError): appearance.bind_generation(request,self.manifest,['object'])
        request['prompt']='Ignore source details'
        self.assertFalse(assets.generation_preflight(request,manifest=self.manifest)['valid'])

    def test_generation_still_needs_authorization_and_attempt_budget(self):
        self.invariant()
        request=appearance.bind_generation(self.request(),self.manifest,['object'])
        request['authorization']=''
        self.assertFalse(assets.generation_preflight(request,manifest=self.manifest)['valid'])
        request['authorization']='approved'; request['attempt']=3
        self.assertFalse(assets.generation_preflight(request,manifest=self.manifest)['valid'])

    def test_generation_does_not_claim_source_exact(self):
        with self.assertRaises(ValueError): appearance.bind_generation(self.request(),self.manifest,['object'])

    def test_review_failure_not_overridden_by_numeric_pass(self):
        path=self.write_evidence()
        self.evidence['appearance_review']['records'][0]['status']='FAIL'
        path.write_text(json.dumps(self.evidence),encoding='utf-8')
        self.assertFalse(appearance.audit(self.artifact,self.manifest,self.root/'manifest.json',path)['valid'])

    def test_runner_does_not_autocertify_legacy_hybrid(self):
        m=dict(editing_policy='hybrid',source_inventory=[])
        path=self.root/'manifest.json'; path.write_text(json.dumps(m),encoding='utf-8')
        result=appearance.load('run-quality-checks').run_checks(self.artifact,path)
        check=next(c for c in result['checks'] if c['check']=='appearance_fidelity')
        self.assertEqual(check['status'],'NOT_VERIFIED')
        self.assertIsNone(result['delivery_ready'])

    def test_asset_binding_coverage_and_legacy_request(self):
        from test_component_assets import Components
        fixture=Components()
        fixture.setUp(); self.addCleanup(fixture.doCleanups)
        self.invariant()
        contract=copy.deepcopy(self.contract)
        contract['source_size']=[200,200]
        contract['source_sha256']=sha(fixture.root/'asset.png')
        contract['items'][0]['id']='sphere-source'
        manifest=fixture.manifest
        manifest['appearance_fidelity']=contract
        asset=manifest['component_assets'][0]
        request=appearance.bind_generation(self.request(),manifest,['sphere-source'])
        request['status']='succeeded'
        asset.update(source_kind='generated',comparison='approved_surrogate',generation=request)
        self.assertEqual(assets.validate_contract(manifest),[])
        self.assertTrue(appearance.load('validate-visual-manifest').validate_manifest(manifest)['valid'])
        extra=copy.deepcopy(contract['items'][0]); extra['id']='label-source'; contract['items'].append(extra)
        asset['generation']=appearance.bind_generation(self.request(),manifest,['sphere-source','label-source'])
        asset['generation']['status']='succeeded'
        self.assertTrue(any('exactly this asset' in e for e in assets.validate_contract(manifest)))
        self.assertTrue(assets.generation_preflight(self.request())['valid'])

    def test_nonfinite_tolerance_and_boolean_flags(self):
        for value in (float('nan'),float('inf'),True,-.01,1):
            self.probe['tolerance']=value
            self.assertTrue(appearance.validate_contract(self.manifest))
        self.probe['tolerance']=.05
        self.manifest['source_inventory'][0]['appearance_sensitive']='true'
        self.assertTrue(appearance.validate_contract(self.manifest))

    def test_embedded_icc_unverified_not_silent_color_conversion(self):
        self.image.save(self.source,icc_profile=b'synthetic-profile-is-not-certified-srgb')
        self.contract['source_sha256']=sha(self.source)
        result=self.audit()
        self.assertTrue(result['unverified'])
        self.assertEqual(result['requirements'][0]['measurement'],'NOT_VERIFIED')

    def test_unknown_review_requirement_rejected(self):
        path=self.write_evidence()
        self.evidence['appearance_review']['records'][0]['requirement_id']='not-in-contract'
        path.write_text(json.dumps(self.evidence),encoding='utf-8')
        self.assertFalse(appearance.audit(self.artifact,self.manifest,self.root/'manifest.json',path)['valid'])

    def test_slide_scope_cannot_inherit_first_slide_result(self):
        self.evidence['slide']=2
        self.assertTrue(self.audit()['unverified'])


if __name__ == '__main__':
    unittest.main()
