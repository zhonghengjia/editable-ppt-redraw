import copy
import hashlib
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from native_components import add_native_component, add_manifest_component, validate_contract
from native_paint import normalize_paint, paint_xml
from appearance_fidelity import sample_paint, verify_sampled_paint
from pptx import Presentation
from pptx.util import Inches
from PIL import Image


def gradient(a='F9E2E0', b='B66476', kind='path'):
    result = {'kind': kind, 'stops': [{'position': 0, 'color': a},
                                     {'position': .55, 'color': 'E7A4AC'},
                                     {'position': 1, 'color': b}]}
    result.update({'focus': [.3, .25]} if kind == 'path' else {'angle': 90})
    return result


def part(pid, path, fill, role='surface', **extra):
    return dict(id=pid, role=role, observation='Synthetic capability fixture, not observed biology',
                d=path, fill=fill, **extra)


def fixture():
    body = 'M 8 46 C 0 12 30 0 55 9 C 91 0 110 28 96 56 C 111 83 77 108 46 95 C 15 107 0 75 8 46 Z'
    nucleus = 'M 36 48 C 28 28 61 25 70 42 C 87 66 62 82 44 70 C 32 66 30 55 36 48 Z'
    def cell(pid, offset, scale):
        return dict(id=pid, role='object', observation='Synthetic grouped object', translate=offset,
                    scale=scale, children=[part(pid+'-body', body, gradient(), stroke='9F6272', stroke_width=.8),
                                          part(pid+'-nucleus', nucleus, gradient('C594B4', '693759'), role='inner part'),
                                          part(pid+'-spot', 'M 49 50 C 49 44 60 44 60 51 C 60 59 49 58 49 50 Z', '793856')])
    cells = dict(id='cells', output_name='cells', viewbox=[0, 0, 210, 155],
                 parts=[cell('back', [65, 0], .9), cell('left', [0, 50], .9), cell('front', [90, 50], 1)])
    tube = dict(id='tube', output_name='tube', viewbox=[0, 0, 180, 150], parts=[
        dict(id='surface', role='host with attached details', observation='Synthetic bent surface and attached band',
             children=[part('host', 'M 10 30 C 35 20 48 74 88 65 C 110 60 122 20 160 25 L 166 58 C 134 56 128 98 88 100 C 50 105 30 57 14 65 Z',
                            gradient('D5ECF5', '476C99', kind='linear')),
                       part('band', 'M 69 63 C 74 65 81 66 88 65 L 88 100 C 80 101 74 100 69 99 Z',
                            gradient('96CFDF', '195069', kind='linear'))])])
    surface = dict(id='surface', output_name='device', viewbox=[0, 0, 150, 150], parts=[
        part('shell', 'M 15 20 L 130 20 L 130 115 L 15 115 Z', gradient('E6EBF1', '586B80', kind='linear')),
        part('inset', 'M 35 45 L 110 45 L 110 80 L 35 80 Z', '273F58'),
        part('flat-mark', 'M 43 53 L 70 53 L 70 60 L 43 60 Z', 'F0F0F0')])
    return [cells, tube, surface]


def build_probe(output):
    """Real generated PPTX fixture consumed by application round-trip qualification."""
    deck = Presentation(); deck.slide_width = Inches(12); deck.slide_height = Inches(4)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i, component in enumerate(fixture()):
        add_native_component(slide, component, .25+i*4, .4, 3.5, 3.2)
    deck.save(output)


class NativeComponentTests(unittest.TestCase):
    def test_declared_missing_contracts_rejected_before_slide_mutation(self):
        component=fixture()[0]
        for flag in ('appearance_sensitive','fidelity_sensitive','structure_sensitive'):
            manifest={'mode':'faithful','source_inventory':[{'id':component['id'],
                'representation':'native_composite',flag:True}],'native_components':[component]}
            prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
            before=slide._element.xml
            with self.subTest(flag=flag),self.assertRaisesRegex(ValueError,'required'):
                add_manifest_component(slide,manifest,component['id'],0,0,5,5,base_dir='.')
            self.assertEqual(slide._element.xml,before)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.folder = Path(self.temp.name)
        self.deck = Presentation(); self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def tearDown(self):
        self.temp.cleanup()

    def source(self, alpha=255):
        image = Image.new('RGBA', (30, 10), (250, 100, 80, alpha))
        image.paste((70, 130, 200, alpha), (10, 0, 20, 10))
        image.paste((20, 40, 60, alpha), (20, 0, 30, 10))
        target = self.folder/'source.png'; image.save(target)
        return target

    def recipe(self, path):
        return dict(source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(), kind='linear',
                    angle=0, positions=[0, .5, 1], patches=[[0,0,10,10], [10,0,10,10], [20,0,10,10]])

    def test_gradient_is_native_and_round_trips(self):
        output = self.folder/'probe.pptx'; build_probe(output)
        loaded = Presentation(output)
        self.assertEqual(len(loaded.slides[0].shapes), 3)
        self.assertEqual(len(loaded.slides[0].shapes[0].shapes), 3)
        with ZipFile(output) as z:
            xml = z.read('ppt/slides/slide1.xml').decode()
            self.assertIn('<a:gradFill', xml); self.assertIn('<a:lin', xml); self.assertIn('<a:path path="circle"', xml)
            self.assertNotIn('<p:pic', xml)
            self.assertFalse(any('/media/' in name for name in z.namelist()))
        ids = [el.get('id') for el in loaded.slides[0]._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')]
        self.assertEqual(len(ids), len(set(ids)))

    def test_preserve_group_hierarchy_and_paint_order(self):
        result = add_native_component(self.slide, fixture()[0], 0, 0, 4, 3)
        root = result['group']
        self.assertEqual([s.name for s in root.shapes], ['cells/back', 'cells/left', 'cells/front'])
        self.assertEqual([s.name for s in root.shapes[0].shapes], ['cells/back/back-body', 'cells/back/back-nucleus', 'cells/back/back-spot'])
        self.assertEqual(len(result['element_map']), 9)

    def test_parent_move_preserves_local_geometry(self):
        root = add_native_component(self.slide, fixture()[0], 0, 0, 4, 3)['group']
        children_before = [s._element.xml for s in root.shapes]
        root.left += Inches(1)
        self.assertEqual(children_before, [s._element.xml for s in root.shapes])

    def test_color_change_does_not_change_path_or_placement(self):
        component = fixture()[0]
        add_native_component(self.slide, component, 0, 0, 4, 3)
        before = self.slide._element.xpath('.//a:custGeom')
        altered = copy.deepcopy(component)
        altered['parts'][0]['children'][0]['fill']['stops'][1]['color'] = 'AABBCC'
        other = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        add_native_component(other, altered, 0, 0, 4, 3)
        self.assertEqual([e.xml for e in before], [e.xml for e in other._element.xpath('.//a:custGeom')])
        self.assertEqual([e.xml for e in self.slide._element.xpath('.//a:xfrm')], [e.xml for e in other._element.xpath('.//a:xfrm')])

    def test_invalid_last_part_is_atomic(self):
        for invalid in ({'kind':'mesh'}, gradient(kind='linear') | {'angle':float('nan')},
                        gradient() | {'focus':[2, 0]}, {'kind':'linear', 'angle':0, 'stops':[]}):
            component = fixture()[0]
            component['parts'][-1]['children'][-1]['fill'] = invalid
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                add_native_component(self.slide, component, 0, 0, 4, 3)
            self.assertEqual(len(self.slide.shapes), 0)

    def test_bad_geometry_and_duplicate_ids_atomic(self):
        for key, value in [('d','M0 broken'), ('rotate',20), ('id','back-body')]:
            component = fixture()[0]
            component['parts'][-1]['children'][-1][key] = value
            with self.subTest(key=key), self.assertRaises(ValueError):
                add_native_component(self.slide, component, 0, 0, 4, 3)
            self.assertEqual(len(self.slide.shapes), 0)

    def test_source_sampling_reads_actual_colors(self):
        path = self.source(); before = path.read_bytes()
        fill = sample_paint(path, self.recipe(path))
        self.assertEqual([s['color'] for s in fill['stops']], ['FA6450', '4682C8', '14283C'])
        self.assertTrue(verify_sampled_paint(path, fill))
        self.assertEqual(path.read_bytes(), before)

    def test_source_hash_and_transparency_rejected(self):
        path = self.source(alpha=128)
        with self.assertRaisesRegex(ValueError, 'transparent'):
            sample_paint(path, self.recipe(path))
        recipe = self.recipe(path); recipe['source_sha256'] = '0'*64
        with self.assertRaisesRegex(ValueError, 'hash'):
            sample_paint(path, recipe)

    def test_bad_source_patch_and_icc_rejected(self):
        path = self.source(); recipe = self.recipe(path); recipe['patches'][0] = [0, 0, 31, 10]
        with self.assertRaises(ValueError): sample_paint(path, recipe)
        with Image.open(path) as image:
            image.save(self.folder/'icc.png', icc_profile=b'unqualified profile')
        path = self.folder/'icc.png'
        with self.assertRaisesRegex(ValueError, 'ICC'): sample_paint(path, self.recipe(path))

    def test_manifest_consumes_sampled_paint_and_checks_source(self):
        path = self.source(); component = fixture()[2]
        component['parts'][0]['fill'] = sample_paint(path, self.recipe(path))
        manifest = dict(source={'path':path.name}, source_inventory=[{'id':'surface', 'representation':'native_composite'}], native_components=[component])
        self.assertEqual(validate_contract(manifest), [])
        add_manifest_component(self.slide, manifest, 'surface', 0, 0, 4, 3, base_dir=self.folder)
        self.assertEqual(len(self.slide.shapes), 1)
        component['parts'][0]['fill']['stops'][0]['color'] = '000000'
        other = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        with self.assertRaisesRegex(ValueError, 'differs'):
            add_manifest_component(other, manifest, 'surface', 0, 0, 4, 3, base_dir=self.folder)
        self.assertEqual(len(other.shapes), 0)

    def test_manifest_identity_and_unknown_fields(self):
        component = fixture()[0]
        manifest = {'source_inventory':[], 'native_components':[component]}
        self.assertTrue(validate_contract(manifest))
        manifest['source_inventory'] = [{'id':'cells', 'representation':'native_composite'}]
        self.assertFalse(validate_contract(manifest))
        component['parts'][0]['opacity'] = .5
        self.assertTrue(validate_contract(manifest))

    def test_duplicate_output_name_refused(self):
        component = fixture()[0]
        add_native_component(self.slide, component, 0, 0, 4, 3)
        with self.assertRaises(ValueError): add_native_component(self.slide, component, 5, 0, 4, 3)
        self.assertEqual(len(self.slide.shapes), 1)

    def test_compound_hole_keeps_subpaths_with_gradient(self):
        component = dict(id='ring', output_name='ring', viewbox=[0,0,100,100], parts=[
            part('rim', 'M0 0 L100 0 L100 100 L0 100 Z M30 30 L30 70 L70 70 L70 30 Z', gradient())])
        add_native_component(self.slide, component, 0, 0, 2, 2)
        self.assertEqual(len(self.slide._element.xpath('.//a:close')), 2)
        self.assertEqual(len(self.slide._element.xpath('.//a:gradFill')), 1)

    def test_manifest_validator_executes_component_contract(self):
        spec = importlib.util.spec_from_file_location('manifest_check_native', ROOT/'scripts/validate-visual-manifest.py')
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        data = dict(schema_version=1, mode='semantic', canvas={'width':100,'height':100},
                    targets=[{'format':'pptx','path':'out.pptx'}], modules=[],
                    source_inventory=[{'id':'cells','representation':'native_composite'}], native_components=[fixture()[0]])
        # Other fields deliberately minimal; only assert routing to this extension.
        data['native_components'][0]['parts'][0]['children'][0]['fill'] = {'kind':'unknown'}
        errors = module.validate_manifest(data)['errors']
        self.assertTrue(any('native_components:' in e for e in errors), errors)

    def test_flat_sample_produces_solid_not_invented_gradient(self):
        path = self.source()
        fill = sample_paint(path, {'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                                  'kind':'solid', 'patches':[[0,0,10,10]]})
        self.assertEqual(fill['kind'], 'solid')
        self.assertEqual(fill['color'], 'FA6450')
        self.assertTrue(verify_sampled_paint(path, fill))

    def test_gradient_stop_alpha_and_focus_serialization(self):
        value = gradient(); value['stops'][0]['alpha'] = .4
        xml = paint_xml(value, .5)
        self.assertIn('alpha val="20000"', xml)
        self.assertIn('l="30000" t="25000" r="70000" b="75000"', xml)
        self.assertIn('pos="55000"', xml)

    def test_paint_strict_schema_and_precision(self):
        for value in ({'kind':'solid','color':'red'}, {'kind':'none','stops':[]},
                      gradient() | {'stops':[{'position':0,'color':'FFFFFF'}, {'position':.000001,'color':'FFFFFF'}, {'position':1,'color':'000000'}]}):
            with self.assertRaises(ValueError): normalize_paint(value)
        self.assertEqual(normalize_paint('#abcdef')['color'], 'ABCDEF')
        self.assertEqual(paint_xml(None), '<a:noFill/>')


if __name__ == '__main__':
    unittest.main()
