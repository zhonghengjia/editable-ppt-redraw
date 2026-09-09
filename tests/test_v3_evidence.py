"""Behavioral regressions for per-instance evidence and honest QA coverage."""
import copy
import hashlib
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

from test_curve_fidelity import AUDIT, EXTRACT, curve_manifest, profile_points, pptx_with_profile
from test_audit_diagram_grammar import manifest as grammar_manifest, layout_elements

ROOT = Path(__file__).resolve().parents[1]


def module(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / 'scripts' / (name + '.py'))
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


TEXT = module('audit-pptx-editability')
TYPE = module('audit-typography-hierarchy')
SOURCE = module('audit-editable-source')
GRAMMAR = module('audit-diagram-grammar')
RUNNER = module('run-quality-checks')


class EvidenceTests(unittest.TestCase):
    def test_numeric_and_word_boundaries(self):
        for actual in ('N=550', 'N=55.0', 'N=55,000', 'N=55a'):
            self.assertEqual(TEXT.text_occurrences('N = 55', actual), 0)
        self.assertEqual(TEXT.text_occurrences('CAR', 'mCAR'), 0)
        self.assertEqual(TEXT.text_occurrences('N = 55', 'N=55;'), 1)
        self.assertEqual(TEXT.text_occurrences('N = 55', 'Study cohort N = 55'), 1)
        self.assertEqual(TEXT.text_occurrences('中文 e\u0301', '中文é'), 1)

    def test_scoped_inventory_counts(self):
        objects = [{'slide': 1, 'id': '2', 'name': 'A', 'text': 'N=55'},
                   {'slide': 2, 'id': '2', 'name': 'B', 'text': 'N=55'}]
        for item in ({'output_slide': 3}, {'output_name': 'A', 'required_count': 2}, {'output_id': '9'}):
            self.assertFalse(TEXT.audit_text_inventory([dict(text='N=55', **item)], objects)[0]['valid'])
        self.assertTrue(TEXT.audit_text_inventory([{'text': 'N=55', 'required_count': 2}], objects)[0]['valid'])

    def test_text_is_not_joined_across_shapes(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'text.pptx'
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for i, text in enumerate(('abc', 'def')):
                slide.shapes.add_textbox(Inches(1), Inches(i+1), Inches(2), Inches(.5)).text = text
            prs.save(path)
            self.assertEqual(TEXT.audit_pptx(path, .8, ['abcdef'])['totals']['missing_required_texts'], 1)

    def curve_check(self, duplicate=False, flip=False, second_wrong=False, rotation=False, multi_path=False, select=None):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / 'curve.pptx'
            pptx_with_profile(output, profile_points())
            with zipfile.ZipFile(output) as archive:
                slide = archive.read('ppt/slides/slide1.xml').decode()
            from xml.etree import ElementTree as ET
            doc = ET.fromstring(slide)
            ns = AUDIT.NS
            tree = doc.find('.//p:spTree', ns)
            shape = tree.find('p:sp', ns)
            if flip or rotation:
                prop = shape.find('p:spPr', ns)
                prop.insert(0, ET.fromstring('<a:xfrm xmlns:a="'+ns['a']+'" '+('flipH="1"' if flip else 'rot="5400000"')+'><a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/></a:xfrm>'))
            if duplicate or second_wrong:
                other = copy.deepcopy(shape)
                other.find('p:nvSpPr/p:cNvPr', ns).set('id', '3')
                if second_wrong:
                    bad = root / 'bad.pptx'
                    pptx_with_profile(bad, profile_points(False))
                    with zipfile.ZipFile(bad) as archive:
                        bad_root = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
                    other.remove(other.find('p:spPr', ns))
                    other.append(bad_root.find('.//p:spPr', ns))
                tree.append(other)
            if multi_path:
                paths = shape.find('.//a:pathLst', ns)
                paths.append(copy.deepcopy(paths[0]))
            with zipfile.ZipFile(output, 'w') as archive:
                archive.writestr('ppt/slides/slide1.xml', ET.tostring(doc))
            (root/'ridge-source.json').write_text(json.dumps({'points': profile_points()}))
            manifest = curve_manifest()
            spec = manifest['curve_fidelity']['series'][0]
            if second_wrong:
                spec['expected_output_count'] = 2
            if flip:
                spec['peak_x_tolerance'] = .01
            if select is not None:
                spec['output_path_index'] = select
            return AUDIT.audit_curve_fidelity(manifest, root/'manifest.json', output)

    def test_duplicate_names_do_not_collapse_objects(self):
        self.assertFalse(self.curve_check(duplicate=True)['valid'])

    def test_every_matching_curve_is_compared(self):
        report = self.curve_check(second_wrong=True)
        self.assertFalse(report['valid'])
        self.assertEqual(len(report['series']['ridge-1']['instances']), 2)

    def test_flip_applied_before_peak_comparison(self):
        self.assertFalse(self.curve_check(flip=True)['valid'])

    def test_rotation_is_not_certified(self):
        report = self.curve_check(rotation=True)
        self.assertFalse(report['valid'])
        self.assertEqual(report['series']['ridge-1']['instances'][0]['status'], 'NOT_VERIFIED')

    def test_multipath_requires_selector(self):
        self.assertFalse(self.curve_check(multi_path=True)['valid'])
        self.assertTrue(self.curve_check(multi_path=True, select=0)['valid'])

    def test_sparse_trace_has_no_invented_points(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'gaps.png'
            image = Image.new('RGB', (101, 51), 'white')
            draw = ImageDraw.Draw(image)
            draw.line((0, 25, 30, 25), fill='red')
            draw.line((70, 25, 100, 25), fill='red')
            image.save(path)
            result = EXTRACT.extract_curve_trace(path, (0, 0, 101, 51), [(255, 0, 0)], mode='centerline', tolerance=1)
            self.assertEqual(result['status'], 'fail')
            self.assertEqual(result['points'], [])

    def test_trace_keeps_crop_height_and_small_gap_evidence(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'height.png'
            image = Image.new('RGB', (101, 51), 'white')
            ImageDraw.Draw(image).rectangle((0, 25, 100, 50), fill='red')
            ImageDraw.Draw(image).rectangle((49, 0, 50, 50), fill='white')
            image.save(path)
            result = EXTRACT.extract_curve_trace(path, (0, 0, 101, 51), [(255, 0, 0)], tolerance=1)
            self.assertEqual(result['status'], 'pass')
            self.assertAlmostEqual(max(p[1] for p in result['points']), .5)
            self.assertEqual(result['trace_quality']['interpolated_columns'], [49, 50])

    def typography(self, runs, exceptions=None):
        spec = {'target_ratio': 1., 'output_name_regex': '^body$'}
        if exceptions:
            spec['run_exceptions'] = exceptions
        manifest = {'typography_hierarchy': {'baseline_role': 'body', 'roles': {'body': spec}}}
        element = {'name': 'body', 'text': 'body', 'resolvedFontSize': 20, 'runs': runs}
        return TYPE.audit_typography(manifest, [{'elements': [element]}])

    def test_unexplained_large_run_fails(self):
        self.assertFalse(self.typography([{'fontSize': s, 'text': 'x'} for s in [20, 20, 80]])['valid'])

    def test_scientific_script_is_not_size_drift(self):
        report = self.typography([{'fontSize': 20, 'text': 'm'}, {'fontSize': 12, 'text': '2', 'baseline': 30000}])
        self.assertEqual(report['status'], 'PASS')

    def test_explicit_run_exception_has_bounds(self):
        runs = [{'fontSize': 20, 'text': 'x'}, {'fontSize': 80, 'text': 'A'}]
        exceptions = [{'text_regex': 'A', 'reason': 'source panel letter', 'min_size_ratio': 4, 'max_size_ratio': 4}]
        self.assertTrue(self.typography(runs, exceptions)['valid'])
        runs[1]['fontSize'] = 100
        self.assertFalse(self.typography(runs, exceptions)['valid'])

    def test_missing_run_evidence_is_not_verified(self):
        self.assertEqual(self.typography([])['status'], 'NOT_VERIFIED')

    def test_drawio_scope_and_explicit_free_ends(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'pages.drawio'
            model = '<mxGraphModel><root><mxCell id="0"/><mxCell id="1" vertex="1"/><mxCell id="e" edge="1" source="1"><mxGeometry relative="1"><mxPoint as="targetPoint" x="10" y="10"/></mxGeometry></mxCell></root></mxGraphModel>'
            path.write_text('<mxfile><diagram>'+model+'</diagram><diagram>'+model+'</diagram></mxfile>')
            report = SOURCE.audit_drawio(path)
            self.assertTrue(report['valid'], report)
            self.assertEqual(report['risks'], [])
            path.write_text(model.replace('source="1"', 'source="missing"'))
            self.assertFalse(SOURCE.audit_drawio(path)['valid'])

    def test_excalidraw_dangling_binding(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'scene.excalidraw'
            scene = {'elements': [{'id': 'a', 'type': 'arrow', 'startBinding': {'elementId': 'missing'}}]}
            path.write_text(json.dumps(scene))
            self.assertFalse(SOURCE.audit_excalidraw(path)['valid'])
            scene['elements'][0]['startBinding'] = None
            path.write_text(json.dumps(scene))
            self.assertTrue(SOURCE.audit_excalidraw(path)['valid'])

    def test_named_lines_are_not_semantic_proof(self):
        report = GRAMMAR.audit_grammar(grammar_manifest(), [{'elements': layout_elements()}])
        self.assertTrue(report['valid'])  # legacy name/geometry result remains scoped
        self.assertEqual(report['status'], 'NOT_VERIFIED')

    def test_wrong_and_correct_native_endpoints(self):
        manifest = grammar_manifest()
        role = manifest['diagram_grammar']['edge_roles']['cohort-to-exclusion']
        role['endpoint_binding'] = {'source_output_name': 'cohort-box', 'target_output_name': 'exclusion-box'}
        elements = layout_elements()[:3]
        elements[2].update(source_name='wrong', target_name='exclusion-box')
        self.assertFalse(GRAMMAR.audit_grammar(manifest, [{'elements': elements}])['valid'])
        elements[2]['source_name'] = 'cohort-box'
        self.assertEqual(GRAMMAR.audit_grammar(manifest, [{'elements': elements}])['status'], 'PASS')

    def test_directed_route_cannot_reverse_or_gap(self):
        manifest = grammar_manifest()
        role = manifest['diagram_grammar']['edge_roles']['cohort-to-exclusion']
        elements = layout_elements()
        elements[0].update(geometry='rect', bbox=[0, 0, 1, 1])
        elements[1].update(bbox=[2, 2, 1, 1])
        elements[2]['endpoints'] = [[1, .5], [2.5, .5]]
        elements[3]['endpoints'] = [[2.5, .5], [2.5, 2]]
        role['endpoint_binding'] = {'source_output_name': 'cohort-box', 'target_output_name': 'exclusion-box', 'ordered_output_names': [e['name'] for e in elements[2:]]}
        self.assertEqual(GRAMMAR.audit_grammar(manifest, [{'elements': elements}])['status'], 'PASS')
        elements[3]['endpoints'].reverse()
        self.assertFalse(GRAMMAR.audit_grammar(manifest, [{'elements': elements}])['valid'])

    def test_runner_fast_artifact_does_not_fake_visual_signoff(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'shape.svg'
            path.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="5" height="5"/></svg>')
            report = RUNNER.run_checks(path, profile='fast')
            self.assertTrue(report['automation_passed'], report)
            self.assertIsNone(report['delivery_ready'])
            self.assertEqual(report['checks'][-1]['status'], 'NOT_VERIFIED')
            self.assertFalse(RUNNER.run_checks(path, profile='dense')['automation_passed'])

    def test_runner_missing_file_fails(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertFalse(RUNNER.run_checks(Path(folder)/'missing.pptx')['automation_passed'])

    def test_runner_invalid_contract_is_reported_not_crashed(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'shape.svg'
            path.write_text('<svg viewBox="0 0 10 10"><rect width="5" height="5"/></svg>')
            contract = Path(folder) / 'manifest.json'
            contract.write_text(json.dumps({'diagram_grammar': 'invalid'}))
            report = RUNNER.run_checks(path, contract)
            self.assertFalse(report['automation_passed'])

    def test_axis_group_transform_and_flip(self):
        from xml.etree import ElementTree as ET
        node = ET.fromstring('<a:xfrm xmlns:a="'+AUDIT.NS['a']+'" flipH="1"><a:off x="100" y="200"/><a:ext cx="400" cy="200"/><a:chOff x="10" y="20"/><a:chExt cx="100" cy="50"/></a:xfrm>')
        sx, sy, tx, ty = AUDIT._axis_transform(node, group=True)
        self.assertEqual((sx*10+tx, sy*20+ty), (500, 200))
        self.assertEqual((sx*110+tx, sy*70+ty), (100, 400))

    def test_runner_does_not_claim_mermaid_parsed(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.mmd'
            path.write_text('flowchart TD\n A --> B')
            report = RUNNER.run_checks(path)
            self.assertFalse(report['automation_passed'])
            self.assertTrue(any(c['check'] == 'native_parser' and c['status'] == 'NOT_VERIFIED' for c in report['checks']))

    def test_real_pptx_runner_end_to_end_and_stale_layout(self):
        from test_validate_visual_manifest import valid_manifest
        from native_vectors import add_rect_link
        from pptx.enum.shapes import MSO_SHAPE
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / 'complete.pptx'
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            nodes = []
            for name, text, x in [('left-box', '稳定', 1), ('right-box', '不稳定', 4)]:
                node = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1), Inches(2), Inches(1))
                node.name, node.text = name, text
                node.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
                nodes.append(node)
            add_rect_link(slide, *nodes, source_port='right', target_port='left', name='flow-left-to-right-link')
            prs.save(output)
            reopened = Presentation(output)
            elements = [{'name': shape.name, 'geometry': 'rect',
                         'bbox': [value/914400 for value in (shape.left, shape.top, shape.width, shape.height)]}
                        for shape in list(reopened.slides[0].shapes)[:2]]
            connection = module('audit-pptx-connections').audit(output)['connectors'][0]
            elements.append({'name': connection['name'], 'geometry': 'line',
                             'source_name': connection['source'], 'target_name': connection['target']})
            manifest = valid_manifest()
            role = manifest['diagram_grammar']['edge_roles']['left-to-right']
            role['endpoint_binding'] = {'source_output_name': 'left-box', 'target_output_name': 'right-box'}
            for item, name in zip(manifest['source_inventory'], ('left-box', 'right-box')):
                item.update(output_slide=1, output_name=name, required_count=1)
            manifest['typography_hierarchy'] = {
                'basis': 'source_observed', 'measurement': 'resolved_font_size', 'baseline_role': 'body',
                'default_tolerance': .1, 'default_max_intra_role_spread': .1,
                'roles': {'body': {'target_ratio': 1., 'output_name_regex': '^(left|right)-box$'}}}
            mf, lf, ef = root/'manifest.json', root/'reopened.layout.json', root/'connections.json'
            mf.write_text(json.dumps(manifest))
            doc = {'artifact_sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'elements': elements}
            lf.write_text(json.dumps(doc))
            ef.write_text(json.dumps([{'slide': 1, 'name': 'flow-left-to-right-link', 'source': 'left-box',
                                      'target': 'right-box', 'start_arrow': 'none', 'end_arrow': 'triangle'}]))
            report = RUNNER.run_checks(output, mf, lf, ef)
            self.assertFalse(report['automation_passed'])
            checks = {c['check']:c for c in report['checks']}
            self.assertEqual(checks['text_clearance']['status'], 'NOT_VERIFIED')
            self.assertTrue(all(c['status'] in ('PASS','NOT_APPLICABLE') for c in report['checks']
                                if c['required_for_automation'] and c['check'] != 'text_clearance'))
            doc['artifact_sha256'] = 'stale'
            lf.write_text(json.dumps(doc))
            stale = RUNNER.run_checks(output, mf, lf, ef)
            self.assertFalse(stale['automation_passed'])
            self.assertEqual(next(c for c in stale['checks'] if c['check']=='layout_provenance')['status'], 'NOT_VERIFIED')


if __name__ == '__main__':
    unittest.main()
