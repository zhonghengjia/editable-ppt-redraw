import copy
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import sys
import unittest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0,str(SCRIPTS))
def load(name):
    spec=importlib.util.spec_from_file_location(name,SCRIPTS/(name+'.py'))
    m=importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m
assets=load('component_assets')

def image_blob(color='blue',size=200,margin=20,mode='RGBA'):
    im=Image.new(mode,(size,size),(0,0,0,0) if mode=='RGBA' else 'white')
    ImageDraw.Draw(im).ellipse((margin,margin,size-margin-1,size-margin-1),fill=color)
    f=io.BytesIO(); im.save(f,format='PNG'); return f.getvalue()

def request():
    return dict(provider='builtin_imagegen',authorization='User requested an isolated synthetic illustration',
                attempt=1,prompt='One object, transparent, no text, no axes, no arrows',subject='synthetic sphere',
                style='flat blue',invariants=['one closed circle'],references=[],status='succeeded')

class Components(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name); self.blob=image_blob(); (self.root/'asset.png').write_bytes(self.blob)
        self.manifest=dict(schema_version=1,mode='faithful',execution_profile='standard',
            source=dict(path='asset.png',width=200,height=200),canvas=dict(width=200,height=200),
            targets=[dict(format='pptx',path='deck.pptx')],modules=[dict(id='panel',bbox=[0,0,200,200])],
            editing_policy='hybrid',hybrid_authorization='User approves object-level image component editing',
            component_assets=[dict(asset_id='a',path='asset.png',sha256=assets.sha(self.blob),source_kind='local',
                origin='Synthetic test fixture; not experimental evidence',authorization='Synthetic fixture approved',
                editing_unit='one independent sphere',extent='object',content_class='illustration',
                contains_native_required=False,requires_alpha=True,min_visible_pixels=50,min_dpi=100,
                invariants=['one closed circle'],comparison='source_exact')],
            component_instances=[dict(instance_id='i',asset_id='a',slide=1,output_name='sphere',bbox_inches=[1,1,1,1],
                crop=[0,0,0,0],rotation=0,placement_tolerance_inches=.001,
                anchors=[dict(id='right',uv=[.9,.5],expected_inches=[1.9,1.5])])],
            source_inventory=[dict(id='sphere-source',module='panel',role='icon',representation='component_raster',component_instance='i'),
                dict(id='label-source',module='panel',role='text',text='Label',representation='native_primitive',output_name='label')])
        self.path=self.root/'deck.pptx'; self.make()

    def make(self, blob=None, grouped=False, extra=False, rotation=0, crop=0):
        prs=Presentation(); s=prs.slides.add_slide(prs.slide_layouts[6]); parent=s.shapes.add_group_shape().shapes if grouped else s.shapes
        p=parent.add_picture(io.BytesIO(blob or self.blob),Inches(1),Inches(1),Inches(1),Inches(1)); p.name='sphere'; p.rotation=rotation; p.crop_left=crop
        t=s.shapes.add_textbox(Inches(3),Inches(1),Inches(2),Inches(.5)); t.name='label'; t.text='Label'
        if extra:
            s.shapes.add_picture(io.BytesIO(self.blob),Inches(4),Inches(3),Inches(1),Inches(1)).name='unregistered'
        prs.save(self.path)

    def audit(self):
        return assets.audit(self.path,self.manifest,self.root/'manifest.json')

    def test_valid_contract_and_actual_media(self):
        self.assertEqual(assets.validate_contract(self.manifest),[])
        self.assertTrue(load('validate-visual-manifest').validate_manifest(self.manifest)['valid'])
        r=self.audit(); self.assertEqual(r['errors'],[]); self.assertEqual(r['unverified'],[])
        self.assertEqual(r['instances'][0]['embedded_sha256'],assets.sha(self.blob))

    def test_legacy_native_unchanged(self):
        m={'schema_version':1,'source_inventory':[]}
        self.assertEqual(assets.validate_contract(m),[])

    def test_requires_hybrid_permission(self):
        for key in ('editing_policy','hybrid_authorization'):
            m=copy.deepcopy(self.manifest); del m[key]
            self.assertTrue(assets.validate_contract(m))

    def test_native_roles_never_flatten(self):
        for role in ('text','chart','connector','table'):
            m=copy.deepcopy(self.manifest); m['source_inventory'][0]['role']=role
            self.assertTrue(assets.validate_contract(m))

    def test_no_fake_native_or_multiple_items_per_picture(self):
        self.manifest['source_inventory'][0]['representation']='native_composite'
        self.assertTrue(assets.validate_contract(self.manifest))
        self.manifest['source_inventory'][0]['representation']='component_raster'
        self.manifest['source_inventory'].append(copy.deepcopy(self.manifest['source_inventory'][0]))
        self.assertTrue(assets.validate_contract(self.manifest))

    def test_source_permission_separate_from_text(self):
        req=request(); self.assertTrue(assets.generation_preflight(req)['valid'])
        req['references']=[dict(path='image.png',sha256='a'*64,role='structure',transmitted_scope='crop')]
        self.assertFalse(assets.generation_preflight(req)['valid'])
        req['references'][0]['authorization']='User approves transfer of this crop'
        self.assertTrue(assets.generation_preflight(req)['valid'])

    def test_generation_failures_attempt_bound_and_no_tool(self):
        self.assertFalse(assets.generation_preflight(request(),False)['valid'])
        req=request(); req['attempt']=3; self.assertFalse(assets.generation_preflight(req)['valid'])
        req['attempt']=2; self.assertFalse(assets.generation_preflight(req)['valid'])
        a=self.manifest['component_assets'][0]; a.update(source_kind='generated',comparison='approved_surrogate',generation=request())
        self.assertEqual(assets.validate_contract(self.manifest),[])
        a['generation']['status']='failed'; self.assertTrue(assets.validate_contract(self.manifest))

    def test_generated_cannot_claim_evidence_or_exact_copy(self):
        a=self.manifest['component_assets'][0]; a.update(source_kind='generated',generation=request())
        self.assertTrue(assets.validate_contract(self.manifest))
        a.update(comparison='approved_surrogate',content_class='evidence'); self.assertTrue(assets.validate_contract(self.manifest))

    def test_hash_drift_on_disk_and_in_package(self):
        (self.root/'asset.png').write_bytes(image_blob('red'))
        self.assertTrue(self.audit()['errors'])
        (self.root/'asset.png').write_bytes(self.blob); self.make(image_blob('red'))
        self.assertTrue(any('embedded media' in e for e in self.audit()['errors']))

    def test_empty_and_unreadable_asset(self):
        for blob in (b'',b'invalid'):
            (self.root/'asset.png').write_bytes(blob)
            self.manifest['component_assets'][0]['sha256']=assets.sha(blob)
            self.assertTrue(self.audit()['errors'])

    def test_opaque_checkerboard_not_transparency(self):
        blob=image_blob(mode='RGB'); (self.root/'asset.png').write_bytes(blob)
        self.manifest['component_assets'][0]['sha256']=assets.sha(blob); self.make(blob)
        self.assertTrue(any('transparency' in e for e in self.audit()['errors']))

    def test_all_transparent_clipped_and_tiny_subject(self):
        raster=load('audit-raster-asset-integrity')
        tiny=image_blob(size=1024,margin=504)
        self.assertTrue(raster.audit_placement(tiny,[1,0,0,1,0,0],[0,0,0,0],50,32)['errors'])
        clear=io.BytesIO(); Image.new('RGBA',(100,100)).save(clear,format='PNG')
        self.assertTrue(raster.audit_placement(clear.getvalue(),[1,0,0,1,0,0],[0,0,0,0],50,32)['errors'])
        self.assertTrue(raster.audit_placement(self.blob,[1,0,0,1,0,0],[.2,0,0,0],50,32)['errors'])

    def test_effective_resolution_and_aspect(self):
        r=load('audit-raster-asset-integrity')
        self.assertTrue(r.audit_placement(self.blob,[10,0,0,10,0,0],[0,0,0,0],100,32)['errors'])
        self.assertTrue(r.audit_placement(self.blob,[2,0,0,1,0,0],[0,0,0,0],50,32)['errors'])
        self.assertEqual(r.audit_placement(self.blob,[1,0,0,1,0,0],[0,0,0,0],100,32)['edge_color_review'],'NOT_VERIFIED')

    def test_recursive_group_readback(self):
        self.make(grouped=True); self.assertEqual(self.audit()['errors'],[])

    def test_nested_rotation_uses_actual_transform(self):
        self.make(rotation=90); i=self.manifest['component_instances'][0]; i['rotation']=90
        i['anchors'][0]['expected_inches']=[1.5,1.9]
        self.assertEqual(self.audit()['errors'],[])
        i['anchors'][0]['expected_inches']=[1.9,1.5]
        self.assertTrue(any('anchor' in e for e in self.audit()['errors']))

    def test_crop_and_unknown_picture_detected(self):
        self.make(crop=.2); self.assertTrue(self.audit()['errors'])
        self.make(extra=True); self.assertTrue(any('unregistered' in e for e in self.audit()['errors']))

    def test_duplicate_asset_instance_and_unknown_ref(self):
        m=copy.deepcopy(self.manifest); m['component_assets'].append(copy.deepcopy(m['component_assets'][0]))
        self.assertTrue(assets.validate_contract(m))
        m=copy.deepcopy(self.manifest); m['component_instances'][0]['asset_id']='missing'
        self.assertTrue(assets.validate_contract(m))

    def test_shared_asset_independent_instances(self):
        prs=Presentation(self.path); p=prs.slides[0].shapes.add_picture(io.BytesIO(self.blob),Inches(4),Inches(3),Inches(1),Inches(1)); p.name='second'; prs.save(self.path)
        i=copy.deepcopy(self.manifest['component_instances'][0]); i.update(instance_id='j',output_name='second',bbox_inches=[4,3,1,1],anchors=[])
        self.manifest['component_instances'].append(i)
        self.manifest['source_inventory'].append(dict(id='second-source',module='panel',role='icon',representation='component_raster',component_instance='j'))
        self.assertEqual(self.audit()['errors'],[])

    def test_replacement_preserves_unrelated_objects(self):
        before=self.root/'before.pptx'; before.write_bytes(self.path.read_bytes())
        new=image_blob('red'); self.make(new)
        self.assertTrue(assets.verify_replacement(before,self.path,1,'sphere',assets.sha(new))['valid'])
        prs=Presentation(self.path); prs.slides[0].shapes[1].text='Altered'; prs.save(self.path)
        self.assertFalse(assets.verify_replacement(before,self.path,1,'sphere',assets.sha(new))['valid'])

    def test_runner_integrates_and_keeps_manual_unverified(self):
        path=self.root/'manifest.json'; path.write_text(json.dumps(self.manifest))
        report=load('run-quality-checks').run_checks(self.path,path)
        check=next(c for c in report['checks'] if c['check']=='editability')
        self.assertEqual(check['status'],'PASS',check)
        manual=next(c for c in report['checks'] if c['check']=='component_visual_review')
        self.assertEqual(manual['status'],'NOT_VERIFIED')
        self.assertIsNone(report['delivery_ready'])

    def test_malformed_contract_rejected_without_crash(self):
        for bad in (None,[],True,{'provider':[]},'bad'):
            m=copy.deepcopy(self.manifest); m['component_assets'][0].update(source_kind='generated',generation=bad)
            self.assertTrue(assets.validate_contract(m))
        m=copy.deepcopy(self.manifest); m['component_instances'][0]['asset_id']=[]
        self.assertTrue(assets.validate_contract(m))

    def test_grouped_whole_image_is_not_editable_pass(self):
        prs=Presentation(); prs.slide_width=Inches(5); prs.slide_height=Inches(5)
        s=prs.slides.add_slide(prs.slide_layouts[6]); g=s.shapes.add_group_shape()
        g.shapes.add_picture(io.BytesIO(self.blob),0,0,Inches(5),Inches(5)).name='sphere'
        t=s.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(.3)); t.name='label'; t.text='Label'; prs.save(self.path)
        i=self.manifest['component_instances'][0]; i.update(bbox_inches=[0,0,5,5],anchors=[])
        self.assertTrue(any('whole-composition' in e for e in self.audit()['errors']))

    def test_combined_tiles_remain_unverified(self):
        prs=Presentation(); prs.slide_width=Inches(2); prs.slide_height=Inches(2)
        s=prs.slides.add_slide(prs.slide_layouts[6]); self.manifest['component_instances']=[]; self.manifest['source_inventory']=[]
        for j,(x,y) in enumerate(((0,0),(1,0),(0,1),(1,1))):
            s.shapes.add_picture(io.BytesIO(self.blob),Inches(x),Inches(y),Inches(1),Inches(1)).name=f'tile-{j}'
            self.manifest['component_instances'].append(dict(instance_id=f'i{j}',asset_id='a',slide=1,output_name=f'tile-{j}',bbox_inches=[x,y,1,1],crop=[0,0,0,0],rotation=0,placement_tolerance_inches=.001,anchors=[]))
            self.manifest['source_inventory'].append(dict(id=f's{j}',module='panel',role='icon',representation='component_raster',component_instance=f'i{j}'))
        prs.save(self.path)
        self.assertTrue(any('possible tiled' in e for e in self.audit()['unverified']))

    def test_group_scale_and_rotation(self):
        self.make(grouped=True)
        prs=Presentation(self.path); group=prs.slides[0].shapes[0]
        group.width=Inches(2); group.height=Inches(2); group.rotation=90; prs.save(self.path)
        i=self.manifest['component_instances'][0]; i.update(bbox_inches=[1,1,2,2],rotation=90,anchors=[])
        self.assertEqual(self.audit()['errors'],[])

    def test_flip_and_effects_are_unverified(self):
        from pptx.oxml.xmlchemy import OxmlElement
        prs=Presentation(self.path); pic=prs.slides[0].shapes[0]; pic._element.spPr.xfrm.set('flipH','1'); prs.save(self.path)
        self.assertTrue(self.audit()['unverified'])
        self.make(); prs=Presentation(self.path); pic=prs.slides[0].shapes[0]
        pic._element.spPr.append(OxmlElement('a:effectLst')); prs.save(self.path)
        self.assertTrue(self.audit()['unverified'])

    def test_source_crop_bounds_and_hash(self):
        a=self.manifest['component_assets'][0]
        a.update(source_kind='source_crop',source_sha256=a['sha256'],source_bbox=[0,0,500,500])
        self.assertTrue(any('exceeds source' in e for e in self.audit()['errors']))
        a.update(source_bbox=[0,0,200,200],source_sha256='0'*64)
        self.assertTrue(any('origin hash' in e for e in self.audit()['errors']))

    def test_resolved_relationship_identity(self):
        import zipfile
        from xml.etree import ElementTree as ET
        before=self.root/'before.pptx'; before.write_bytes(self.path.read_bytes())
        with zipfile.ZipFile(self.path) as z: parts={n:z.read(n) for n in z.namelist()}
        slide='ppt/slides/slide1.xml'; rel='ppt/slides/_rels/slide1.xml.rels'
        r=ET.fromstring(parts[slide]); blip=r.find('.//a:blip',assets.NS); key='{'+assets.NS['r']+'}embed'; old=blip.get(key);blip.set(key,'newImageRelationship')
        rels=ET.fromstring(parts[rel])
        for node in rels:
            if node.get('Id')==old: node.set('Id','newImageRelationship')
        parts[slide]=ET.tostring(r);parts[rel]=ET.tostring(rels)
        with zipfile.ZipFile(self.path,'w') as z:
            for n,b in parts.items(): z.writestr(n,b)
        self.assertTrue(assets.verify_replacement(before,self.path,1,'sphere',assets.sha(self.blob))['valid'])

if __name__=='__main__': unittest.main()
