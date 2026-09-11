"""Synthetic, source-bound fixtures: never model output or licensed demo art."""
import copy
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import component_assets as ca
import raster_components as rc


def blob(image):
    out = io.BytesIO()
    image.save(out, format='PNG')
    return out.getvalue()


class RasterComponents(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        im = Image.new('RGBA', (200, 200))
        d = ImageDraw.Draw(im)
        d.ellipse((20, 20, 180, 180), fill=(36, 110, 190, 255))
        d.ellipse((70, 70, 120, 120), fill=(0, 0, 0, 0))
        d.line((90, 20, 90, 5), fill=(100, 180, 220, 80), width=2)
        self.source = blob(im)
        (self.root / 'source.png').write_bytes(self.source)
        self.m = dict(schema_version=1, mode='faithful', execution_profile='standard',
            source=dict(path='source.png', width=200, height=200), canvas=dict(width=200, height=200),
            targets=[dict(format='pptx', path='out.pptx')], modules=[dict(id='panel', bbox=[0, 0, 200, 200])],
            editing_policy='hybrid', hybrid_authorization='Synthetic fixture: picture editing approved',
            component_assets=[dict(asset_id='a', path='prepared.png', source_kind='source_crop',
                source_sha256=rc.digest(self.source), source_bbox=[0, 0, 200, 200],
                origin='Synthetic ring with a fine branch', authorization='Fixture only',
                editing_unit='one ring', extent='object', content_class='illustration',
                contains_native_required=False, requires_alpha=True, min_visible_pixels=50, min_dpi=100,
                invariants=['real hole and fine branch'], comparison='source_exact',
                preparation=dict(method='preserve-alpha', authorization='Preserve source alpha'))],
            component_instances=[dict(instance_id='i', asset_id='a', slide=1, output_name='ring',
                bbox_inches=[1, 1, 1, 1], crop=[0, 0, 0, 0], rotation=0, placement_tolerance_inches=.001,
                anchors=[dict(id='right', uv=[.9, .5], expected_inches=[1.9, 1.5])])],
            source_inventory=[dict(id='ring-source', module='panel', role='icon',
                representation='component_raster', component_instance='i'),
                dict(id='label-source', module='panel', role='text', text='Label',
                representation='native_primitive', output_name='label')])

    def prepared(self):
        data, manifest, report = rc.prepare_asset(self.m, 'a', self.root)
        (self.root / 'prepared.png').write_bytes(data)
        return data, manifest, report

    def deck(self, manifest, path):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        assets = {a['asset_id']: a for a in manifest['component_assets']}
        for i in manifest['component_instances']:
            slide.shapes.add_picture(io.BytesIO((self.root / assets[i['asset_id']]['path']).read_bytes()),
                *[Inches(x) for x in i['bbox_inches']]).name = i['output_name']
        slide.shapes.add_textbox(Inches(3), Inches(1), Inches(2), Inches(.5)).name = 'label'
        slide.shapes[-1].text = 'Label'
        prs.save(path)

    def supplied(self, mode='straight', matte=(255, 255, 255)):
        alpha = np.zeros((200, 200), dtype=np.uint8)
        alpha[20:180, 20:180] = 128
        alpha[70:120, 70:120] = 0
        rgb = np.zeros((200, 200, 3), dtype=np.uint8) + [51, 153, 204]
        if mode == 'matted':
            rgb = np.rint(rgb * (alpha[:, :, None]/255) + np.array(matte) * (1-alpha[:, :, None]/255))
        source = blob(Image.fromarray(rgb.astype(np.uint8)))
        mask = blob(Image.fromarray(alpha))
        (self.root / 'source.png').write_bytes(source)
        (self.root / 'alpha.png').write_bytes(mask)
        a = self.m['component_assets'][0]
        a.update(source_sha256=rc.digest(source), preparation=dict(method='supplied-alpha',
            authorization='Given synthetic alpha', alpha_path='alpha.png', alpha_sha256=rc.digest(mask), rgb_mode=mode))
        if mode == 'matted':
            a['preparation']['matte_rgb'] = list(matte)
        return alpha

    def test_preserves_texture_hole_branch_and_manifest(self):
        before = copy.deepcopy(self.m)
        data, m, report = self.prepared()
        self.assertEqual(self.m, before)
        self.assertEqual(Image.open(io.BytesIO(data)).tobytes(), Image.open(io.BytesIO(self.source)).tobytes())
        self.assertEqual(m['component_assets'][0]['sha256'], rc.digest(data))
        self.assertFalse(report['alpha_inferred'])
        self.assertEqual(report['visual_fidelity'], 'NOT_VERIFIED')

    def test_supplied_straight_alpha(self):
        alpha = self.supplied()
        data, _, _ = self.prepared()
        out = np.asarray(Image.open(io.BytesIO(data)))
        np.testing.assert_array_equal(out[:, :, 3], alpha)
        np.testing.assert_array_equal(out[30, 30, :3], [51, 153, 204])
        np.testing.assert_array_equal(out[90, 90], [0, 0, 0, 0])

    def test_known_matte_roundtrip_and_foreground(self):
        for matte in ((255, 255, 255), (10, 20, 30)):
            with self.subTest(matte=matte):
                self.supplied('matted', matte)
                data, _, report = self.prepared()
                out = np.asarray(Image.open(io.BytesIO(data)))
                self.assertLessEqual(np.max(np.abs(out[30, 30, :3].astype(int)-[51, 153, 204])), 1)
                self.assertLessEqual(report['visible_roundtrip_max_channel_error'], .51)

    def test_low_alpha_quantization_not_hidden_generation(self):
        self.supplied('matted')
        mask = blob(Image.new('L', (200, 200), 1))
        source = blob(Image.new('RGB', (200, 200), (254, 255, 254)))
        (self.root/'alpha.png').write_bytes(mask); (self.root/'source.png').write_bytes(source)
        a = self.m['component_assets'][0]
        a['source_sha256'] = rc.digest(source); a['preparation']['alpha_sha256'] = rc.digest(mask)
        data, _, report = self.prepared()
        self.assertEqual(Image.open(io.BytesIO(data)).getpixel((50, 50)), (0, 255, 0, 1))
        self.assertFalse(report['hidden_content_generated'])

    def test_wrong_matte_rejected(self):
        self.supplied('matted')
        self.m['component_assets'][0]['preparation']['matte_rgb'] = [0, 0, 0]
        with self.assertRaisesRegex(ValueError, 'incompatible'):
            self.prepared()

    def test_hash_and_mask_size_rejected(self):
        self.supplied()
        (self.root/'alpha.png').write_bytes(blob(Image.new('L', (10, 10), 128)))
        with self.assertRaisesRegex(ValueError, 'hash mismatch'):
            self.prepared()
        self.m['component_assets'][0]['preparation']['alpha_sha256'] = rc.digest((self.root/'alpha.png').read_bytes())
        with self.assertRaisesRegex(ValueError, 'exactly match'):
            self.prepared()

    def test_source_dimensions_and_recipe_drift(self):
        _, m, _ = self.prepared()
        self.m['source']['width'] = 201
        with self.assertRaisesRegex(ValueError, 'dimensions'): self.prepared()
        (self.root/'source.png').write_bytes(blob(Image.new('RGB', (200, 200), 'red')))
        with self.assertRaisesRegex(ValueError, 'hash'):
            rc.verify_preparation(m, 'a', self.root, (self.root/'prepared.png').read_bytes())

    def test_double_alpha_rejected(self):
        self.supplied()
        (self.root/'source.png').write_bytes(self.source)
        self.m['component_assets'][0]['source_sha256'] = rc.digest(self.source)
        with self.assertRaisesRegex(ValueError, 'apply alpha twice'):
            self.prepared()

    def test_clear_and_opaque_rejected(self):
        for color, reason in (((0, 0, 0, 0), 'fully transparent'), ((1, 2, 3, 255), 'transparency absent')):
            data = blob(Image.new('RGBA', (200, 200), color))
            (self.root/'source.png').write_bytes(data)
            self.m['component_assets'][0]['source_sha256'] = rc.digest(data)
            with self.assertRaisesRegex(ValueError, reason):
                self.prepared()

    def test_crop_bounds_integrality_and_authorization(self):
        for box in ([0, 0, 201, 200], [0, 0, 200.0, 200], [-1, 0, 200, 200]):
            self.m['component_assets'][0]['source_bbox'] = box
            with self.assertRaises(ValueError): self.prepared()
        self.m['component_assets'][0]['source_bbox'] = [0, 0, 200, 200]
        self.m['component_assets'][0]['preparation']['authorization'] = ''
        with self.assertRaises(ValueError): self.prepared()

    def test_preparation_schema_and_native_boundary(self):
        for prep in (None, [], {}, {'method': 'erase-white', 'authorization': 'x'},
                     {'method': 'preserve-alpha', 'authorization': 'x', 'threshold': 128}):
            a = copy.deepcopy(self.m['component_assets'][0]); a['preparation'] = prep
            self.assertTrue(rc.validate_preparation(a))
        self.m['editing_policy'] = 'native'
        with self.assertRaises(ValueError): self.prepared()

    def test_readback_checks_recipe_not_just_updated_hash(self):
        data, m, _ = self.prepared()
        self.deck(m, self.root/'before.pptx')
        r = ca.audit(self.root/'before.pptx', m, self.root/'manifest.json')
        self.assertEqual(r['errors'], [], r)
        im = Image.open(io.BytesIO(data)); im.putpixel((30, 30), (255, 0, 0, 255))
        bad = blob(im); (self.root/'prepared.png').write_bytes(bad)
        m['component_assets'][0]['sha256'] = rc.digest(bad)
        self.deck(m, self.root/'after.pptx')
        r = ca.audit(self.root/'after.pptx', m, self.root/'manifest.json')
        self.assertTrue(any('pixels differ' in e for e in r['errors']), r)

    def test_color_mode_icc_orientation_frame_limits(self):
        for key, image, kwargs in (
            ('mode', Image.new('P', (8, 8)), {}),
            ('ICC', Image.new('RGB', (8, 8)), {'icc_profile': b'unqualified'}),
            ('EXIF', Image.new('RGB', (8, 8)), {'exif': Image.Exif()})):
            if key == 'EXIF': kwargs['exif'][274] = 6
            path = self.root/(key+'.png'); image.save(path, **kwargs)
            with self.assertRaises(ValueError): rc.read_image(path)
        path = self.root/'animation.png'
        Image.new('RGBA', (10, 10), 'red').save(path, save_all=True, append_images=[Image.new('RGBA', (10, 10), 'blue')])
        with self.assertRaisesRegex(ValueError, 'frame'): rc.read_image(path)
        mask = self.root/'rgbmask.png'; Image.new('RGB', (10, 10)).save(mask)
        with self.assertRaisesRegex(ValueError, 'alpha requires'): rc.read_image(mask, alpha=True)

    def shared(self, m):
        i = copy.deepcopy(m['component_instances'][0])
        i.update(instance_id='j', output_name='second', anchors=[])
        m['component_instances'].append(i)
        m['source_inventory'].append(dict(id='second-source', module='panel', role='icon',
            representation='component_raster', component_instance='j'))

    def new_asset(self, m, size=(200, 200)):
        image = Image.new('RGBA', size)
        ImageDraw.Draw(image).ellipse((20, 20, size[0]-20, size[1]-20), fill=(210, 40, 80, 128))
        data = blob(image); (self.root/'new.png').write_bytes(data)
        asset = copy.deepcopy(m['component_assets'][0])
        for k in ('preparation', 'source_sha256', 'source_bbox'): asset.pop(k, None)
        asset.update(asset_id='b', path='new.png', sha256=rc.digest(data), source_kind='local')
        return asset

    def test_immutable_shared_replacement_and_actual_pptx(self):
        _, m, _ = self.prepared(); self.shared(m)
        m['component_instances'][1]['bbox_inches'] = [4, 3, 1, 1]
        frozen = copy.deepcopy(m)
        new = self.new_asset(m)
        after = rc.replace_instance_asset(m, 'i', new, self.root)
        self.assertEqual(m, frozen)
        self.assertEqual(after['component_instances'][1], m['component_instances'][1])
        expected = dict(m['component_instances'][0], asset_id='b')
        self.assertEqual(after['component_instances'][0], expected)
        self.assertEqual(len(after['component_assets']), 2)
        self.deck(m, self.root/'before.pptx'); self.deck(after, self.root/'after.pptx')
        result = ca.verify_replacement(self.root/'before.pptx', self.root/'after.pptx', 1, 'ring', new['sha256'])
        self.assertTrue(result['valid'], result)

    def test_replacement_rejects_reused_id_shape_and_hash(self):
        _, m, _ = self.prepared(); new = self.new_asset(m, (200, 100))
        with self.assertRaisesRegex(ValueError, 'aspect ratio'): rc.replace_instance_asset(m, 'i', new, self.root)
        new = self.new_asset(m); new['asset_id'] = 'a'
        with self.assertRaisesRegex(ValueError, 'fresh'): rc.replace_instance_asset(m, 'i', new, self.root)
        new['asset_id'] = 'b'; new['sha256'] = '0'*64
        with self.assertRaisesRegex(ValueError, 'hash'): rc.replace_instance_asset(m, 'i', new, self.root)

    def test_preview_order_and_alpha_against_numeric_reference(self):
        _, m, _ = self.prepared(); self.shared(m)
        new = self.new_asset(m); m = rc.replace_instance_asset(m, 'j', new, self.root)
        image, report = rc.render_components(m, self.root, ['i', 'j'], canvas_inches=[3, 3], dpi=200)
        expected = np.rint(np.array([210, 40, 80])*128/255 + np.array([36, 110, 190])*127/255)
        np.testing.assert_allclose(image.getpixel((330, 300))[:3], expected, atol=1)
        reverse, _ = rc.render_components(m, self.root, ['j', 'i'], canvas_inches=[3, 3], dpi=200)
        self.assertEqual(reverse.getpixel((330, 300)), (36, 110, 190, 255))
        self.assertEqual(report['scope'], 'PARTIAL_COMPONENT_PREVIEW'); self.assertIsNone(report['delivery_ready'])
        self.assertFalse(report['native_objects_rendered'])

    def test_premultiplied_resizing_excludes_hidden_rgb(self):
        _, m, _ = self.prepared()
        array = np.zeros((200, 200, 4), dtype=np.uint8)
        array[:, :, 0] = 255  # Invisible red must not contaminate visible blue edges.
        array[20:180, 20:180] = [0, 0, 255, 255]
        data = blob(Image.fromarray(array)); (self.root/'prepared.png').write_bytes(data)
        m['component_assets'][0]['sha256'] = rc.digest(data)
        image, _ = rc.render_components(m, self.root, ['i'], canvas_inches=[3, 3], dpi=43, background=[0, 0, 0])
        self.assertEqual(np.asarray(image)[:, :, 0].max(), 0)

    def test_preview_refuses_unsupported_or_bad_inputs(self):
        _, m, _ = self.prepared()
        for kwargs in ({'order': ['i', 'i']}, {'order': ['missing']}, {'dpi': 0},
                       {'canvas_inches': [float('nan'), 3]}, {'canvas_inches': [1000, 1000]},
                       {'canvas_inches': [1, 1]}, {'background': [0, 0, 256]}):
            args = dict(order=['i'], canvas_inches=[3, 3]); args.update(kwargs)
            with self.assertRaises(ValueError): rc.render_components(m, self.root, **args)
        for key, value in (('rotation', 90), ('crop', [.1, 0, 0, 0]), ('bbox_inches', [1, 1, 1, .5])):
            altered = copy.deepcopy(m); altered['component_instances'][0][key] = value
            with self.assertRaises(ValueError): rc.render_components(altered, self.root, ['i'], canvas_inches=[3, 3])

    def cli(self, *args):
        return subprocess.run([sys.executable, '-B', str(Path(rc.__file__)), *args], capture_output=True, text=True, cwd=self.root)

    def test_cli_prepare_preview_and_no_overwrite(self):
        path = self.root/'plan.json'; path.write_text(json.dumps(self.m))
        args = ['prepare', '--manifest', str(path), '--asset-id', 'a', '--out-manifest', str(self.root/'ready.json')]
        result = self.cli(*args); self.assertEqual(result.returncode, 0, result.stderr)
        self.assertFalse(json.loads(result.stdout)['alpha_inferred'])
        self.assertNotEqual(self.cli(*args).returncode, 0)
        result = self.cli('preview', '--manifest', str(self.root/'ready.json'), '--order', 'i',
            '--canvas-inches', '3', '3', '--output', str(self.root/'preview.png'))
        self.assertEqual(result.returncode, 0, result.stderr)
        with Image.open(self.root/'preview.png') as image:
            self.assertEqual(image.size, (288, 288))
        self.assertEqual(json.loads(path.read_text()), self.m)

    def test_cli_rejects_target_manifest_collision_and_escape(self):
        path = self.root/'plan.json'
        for target in ('ready.json', '../escaped.png'):
            self.m['component_assets'][0]['path'] = target; path.write_text(json.dumps(self.m))
            result = self.cli('prepare', '--manifest', str(path), '--asset-id', 'a', '--out-manifest', str(self.root/'ready.json'))
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse((self.root/'ready.json').exists())

    def test_cli_replace_keeps_original_and_no_deck_mutation(self):
        _, m, _ = self.prepared(); new = self.new_asset(m)
        path = self.root/'ready.json'; path.write_text(json.dumps(m))
        asset = self.root/'replacement.json'; asset.write_text(json.dumps(new))
        args = ['replace', '--manifest', str(path), '--instance-id', 'i', '--asset-record', str(asset),
                '--out-manifest', str(self.root/'replacement-ready.json')]
        result = self.cli(*args); self.assertEqual(result.returncode, 0, result.stderr)
        self.assertFalse(json.loads(result.stdout)['pptx_regenerated'])
        self.assertEqual(json.loads(path.read_text()), m)
        candidate = json.loads((self.root/'replacement-ready.json').read_text())
        self.assertEqual(candidate['component_instances'][0]['asset_id'], 'b')
        self.assertEqual(len(candidate['component_assets']), 1)
        self.assertNotEqual(self.cli(*args).returncode, 0)


if __name__ == '__main__':
    unittest.main()
