"""General adversarial geometry and actual-file evidence regressions."""
import copy
import importlib.util
import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
import component_fidelity as CF
import native_vectors as NV


def module(name):
    spec = importlib.util.spec_from_file_location(name, ROOT/'scripts'/f'{name}.py')
    value = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(value)
    return value


NODE = os.environ.get('EDITABLE_PPT_NODE') or shutil.which('node')


def manifest(source_hash='0'*64):
    return {'schema_version': 1, 'mode': 'faithful', 'execution_profile': 'standard',
            'source': {'path': 'source.png'}, 'canvas': {'width': 100, 'height': 100},
            'targets': [{'format': 'pptx', 'path': 'final.pptx'}],
            'modules': [{'id': 'panel', 'bbox': [0,0,100,100]}],
            'source_inventory': [{'id': 'symbol', 'module': 'panel', 'role': 'shape',
                                  'representation': 'native_composite', 'fidelity_sensitive': True}],
            'regional_fidelity': {'source_size': [100,100], 'source_sha256': source_hash,
                'regions': [{'id': 'symbol', 'source_bbox': [20,20,60,60], 'features': ['hole','thin stalk','inner band'],
                             'threshold': .1, 'window': 5, 'max_mismatch_ratio': .02,
                             'max_window_ratio': .19, 'rationale': 'Synthetic exact-color geometry; allow isolated pixel noise only'}]}}


class ContractTests(unittest.TestCase):
    def test_manifest_integration(self):
        self.assertTrue(module('validate-visual-manifest').validate_manifest(manifest())['valid'])
        value = manifest()
        del value['regional_fidelity']
        self.assertFalse(module('validate-visual-manifest').validate_manifest(value)['valid'])

    def test_empty_unknown_duplicate_regions(self):
        for change in ('empty','unknown','duplicate','missing'):
            value = manifest()
            regions = value['regional_fidelity']['regions']
            if change == 'empty': regions.clear()
            elif change == 'unknown': regions[0]['id'] = 'wrong'
            elif change == 'duplicate': regions.append(copy.deepcopy(regions[0]))
            else: value['source_inventory'].append(dict(value['source_inventory'][0], id='second'))
            self.assertTrue(CF.validate_contract(value), change)

    def test_bad_thresholds_coordinates_features(self):
        for key, val in [('threshold',float('nan')), ('max_window_ratio',1), ('window',True),
                         ('source_bbox',[0,0,101,2]), ('source_bbox',[0,0,1.5,2]),
                         ('rationale',''), ('features',[])]:
            value = manifest()
            value['regional_fidelity']['regions'][0][key] = val
            self.assertTrue(CF.validate_contract(value), key)
        for size in ([], [True,10], [0,10]):
            value = manifest(); value['regional_fidelity']['source_size'] = size
            self.assertTrue(CF.validate_contract(value))

    def test_redesign_is_not_pixel_faithful(self):
        value = manifest(); value['mode'] = 'redesign'
        self.assertTrue(CF.validate_contract(value))

    def test_missing_evidence_remains_unverified(self):
        report = CF.audit_fidelity('unused', manifest(), 'unused', None)
        self.assertTrue(report['valid'])
        self.assertTrue(report['unverified'])

    def test_multiline_native_text(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)/'text.pptx'
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text = 'ACSL1\nCPT1A'
            prs.save(output)
            audit = module('audit-pptx-editability')
            with zipfile.ZipFile(output) as archive:
                objects = audit.text_objects(archive, ['ppt/slides/slide1.xml'])
            self.assertEqual(objects[0]['text'], 'ACSL1\nCPT1A')
            self.assertTrue(audit.audit_text_inventory([{'text':'ACSL1 CPT1A'}], objects)[0]['valid'])

    def test_vendor_hashes(self):
        lock = json.loads((ROOT/'references/upstream-lock.json').read_text())
        for entry in lock['components']:
            if entry['repo'] not in ('jankovicsandras/imagetracerjs','mapbox/pixelmatch'): continue
            self.assertEqual(CF.digest(ROOT/entry['bundled_path']), entry['source_sha256'])
            self.assertEqual(CF.digest(ROOT/entry['license_path']), entry['license_sha256'])


@unittest.skipUnless(NODE, 'requires existing Node; set EDITABLE_PPT_NODE to test integration')
class ComponentIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.source = Image.new('RGB', (100,100), 'white')
        draw = ImageDraw.Draw(self.source)
        draw.ellipse((25,25,75,65), fill='#2277aa')
        draw.ellipse((40,36,49,45), fill='white')
        draw.rectangle((55,28,58,60), fill='#cc4433')
        draw.rectangle((49,63,50,78), fill='#2277aa')
        self.source.save(self.root/'source.png')
        prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(self.root/'final.pptx')
        self.value = manifest(CF.digest(self.root/'source.png'))

    def tearDown(self): self.temp.cleanup()

    def audit(self, image=None, evidence_change=None, node=NODE):
        (image or self.source).save(self.root/'render.png')
        evidence = {'artifact_sha256':CF.digest(self.root/'final.pptx'),
                    'render_path':'render.png', 'render_sha256':CF.digest(self.root/'render.png'),
                    'renderer':'synthetic test fixture, not a claimed Office render',
                    'rendered_from_final_artifact':True}
        if evidence_change: evidence.update(evidence_change)
        (self.root/'evidence.json').write_text(json.dumps(evidence))
        return CF.audit_fidelity(self.root/'final.pptx', self.value, self.root/'manifest.json', self.root/'evidence.json', node)

    def test_identical_passes_but_no_semantic_signoff(self):
        report = self.audit()
        self.assertTrue(report['valid'], report)
        self.assertEqual(report['regions'][0]['mismatch_pixels'], 0)
        self.assertEqual(report['semantic_review'], 'NOT_VERIFIED')

    def test_inner_band_hole_stalk_and_occlusion_fail(self):
        for defect in ('band','hole','stalk','occlusion'):
            image = self.source.copy(); draw = ImageDraw.Draw(image)
            if defect == 'band': draw.rectangle((55,28,58,60), fill='#2277aa')
            elif defect == 'hole': draw.ellipse((40,36,49,45), fill='#2277aa')
            elif defect == 'stalk': draw.rectangle((49,66,50,78), fill='white')
            else: draw.rectangle((47,34,61,46), fill='#cc4433')
            report = self.audit(image)
            self.assertFalse(report['valid'], defect)
            self.assertGreater(report['regions'][0]['worst_window_ratio'], .19)

    def test_local_window_catches_global_dilution(self):
        image = self.source.copy(); ImageDraw.Draw(image).rectangle((49,70,50,78), fill='white')
        self.value['regional_fidelity']['regions'][0]['source_bbox'] = [0,0,100,100]
        report = self.audit(image)
        self.assertLess(report['regions'][0]['mismatch_ratio'], .02)
        self.assertFalse(report['valid'])

    def test_wrong_hash_and_aspect_rejected(self):
        self.assertFalse(self.audit(evidence_change={'artifact_sha256':'0'*64})['valid'])
        self.assertFalse(self.audit(evidence_change={'render_sha256':'0'*64})['valid'])
        self.assertFalse(self.audit(self.source.resize((110,100)))['valid'])
        self.assertFalse(self.audit(self.source.resize((90,90)))['valid'])
        self.value['regional_fidelity']['source_size'] = [101,100]
        self.assertFalse(self.audit()['valid'])

    def test_missing_runtime_unverified(self):
        report = self.audit(node=str(self.root/'no-node'))
        self.assertTrue(report['unverified'])
        self.assertFalse(report['regions'])

    def test_runner_enforces_render_contract(self):
        self.audit()
        (self.root/'manifest.json').write_text(json.dumps(self.value))
        run = module('run-quality-checks')
        report = run.run_checks(self.root/'final.pptx', self.root/'manifest.json')
        check = next(c for c in report['checks'] if c['check'] == 'regional_fidelity')
        self.assertEqual(check['status'], 'NOT_VERIFIED')
        self.assertFalse(report['automation_passed'])

    def test_trace_deterministic_native_curves_and_holes(self):
        traces = []
        for name in ('first','second'):
            traces.append(CF.trace_component(self.root/'source.png', self.root/f'{name}.svg', [20,20,60,60], 4, NODE))
        self.assertEqual(traces[0]['svg_sha256'], traces[1]['svg_sha256'])
        self.assertFalse(traces[0]['semantics_verified'])
        specs, _, _ = NV.read_vectors(self.root/'first.svg')
        self.assertIn('CC4433', [s['fill'] for s in specs], 'sparse sampling must not erase a rare inner band')
        self.assertTrue(any(sum(c.cmd == 'M' for c in s['commands']) > 1 for s in specs), 'compound holes retained')
        self.assertTrue(any(c.cmd in ('Q','C') for s in specs for c in s['commands']))
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        NV.add_svg_component(slide, self.root/'first.svg', 1,1,3,3,name='traced')
        pptx = self.root/'native.pptx'; prs.save(pptx)
        with zipfile.ZipFile(pptx) as archive:
            xml = archive.read('ppt/slides/slide1.xml').decode()
        self.assertIn('custGeom', xml)
        self.assertIn('cubicBezTo', xml)
        self.assertNotIn('<p:pic>', xml)

    def test_palette_covers_thin_rare_color(self):
        image = Image.new('RGBA', (100,100), 'white')
        ImageDraw.Draw(image).point((1,1), fill='#00aa44')
        palette, basis = CF.source_palette(image,2)
        self.assertIn({'r':0,'g':170,'b':68,'a':255},palette)
        self.assertEqual(basis,'all_observed_rgba_colors')

    def test_trace_refuses_overwrite_invalid_crop_and_colors(self):
        output = self.root/'trace.svg'
        for crop in ([0,0,101,1], [-1,0,4,4], [0,0,0,4]):
            with self.assertRaises(ValueError): CF.trace_component(self.root/'source.png', output, crop, node=NODE)
        for colors in (0,65):
            with self.assertRaises(ValueError): CF.trace_component(self.root/'source.png', output, [20,20,60,60],colors,NODE)
        output.write_text('protected')
        with self.assertRaises(ValueError): CF.trace_component(self.root/'source.png', output,[20,20,60,60],4,NODE)
        self.assertEqual(output.read_text(), 'protected')


if __name__ == '__main__': unittest.main()
