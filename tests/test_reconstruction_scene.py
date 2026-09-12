import io
from pathlib import Path
import sys
import tempfile
import unittest

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from raster_components import digest
from reconstruction_scene import (ReconstructionScene, SceneNode, SourceFrame,
    assemble_pptx, rectangle_draw, svg_draw, text_draw)


class ReconstructionScenes(unittest.TestCase):
    def setUp(self):
        self.frame = SourceFrame((100, 200, 400, 200), (1, 1, 8, 4))
        self.deck = Presentation()
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def node(self, identity, **kwargs):
        return SceneNode(identity, (120, 220, 30, 30), rectangle_draw(fill='3399AA'), **kwargs)

    def test_container_follows_ownership_not_append_order(self):
        child = self.node('protein', parent='container')
        parent = SceneNode('container', (110, 210, 80, 80), rectangle_draw(fill='FFFFFF'), role='container', z=100)
        scene = ReconstructionScene(self.frame, [child, parent])
        report = assemble_pptx(self.slide, scene)
        stream = io.BytesIO()
        self.deck.save(stream)
        reread = Presentation(io.BytesIO(stream.getvalue())).slides[0]
        self.assertEqual([s.name for s in reread.shapes], ['container', 'protein'])
        self.assertEqual(report['paint_order'], ['container', 'protein'])

    def test_uniform_frame_nonzero_origin(self):
        f = SourceFrame((100, 200, 400, 200), (1, 1, 8, 8))
        self.assertEqual(f.place((100, 200, 400, 200)), (1, 3, 8, 4))
        self.assertEqual(f.place((200, 250, 100, 50)), (3, 4, 2, 1))

    def test_explicit_order_and_stability(self):
        scene = ReconstructionScene(self.frame, [self.node('a', behind=('b',)), self.node('b', z=-2), self.node('c')])
        self.assertEqual([n.id for n in scene.ordered()], ['a', 'b', 'c'])

    def test_bad_topology_fails_before_emission(self):
        cases = [[self.node('a', behind=('b',)), self.node('b', behind=('a',))],
                 [self.node('a'), self.node('a')], [self.node('a', parent='missing')],
                 [self.node('a', behind=('missing',))],
                 [self.node('edge', role='relation')],
                 [self.node('edge', role='relation', relation=('missing', 'missing'))]]
        for nodes in cases:
            with self.subTest(nodes=[n.id for n in nodes]), self.assertRaises(ValueError):
                assemble_pptx(self.slide, ReconstructionScene(self.frame, nodes))
            self.assertEqual(len(self.slide.shapes), 0)

    def test_source_curve_controls_reach_native_xml_and_receipt(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'curve.svg'
            path.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50"><path id="route" d="M0 25 C90 0 10 50 100 25" fill="none" stroke="#168040" stroke-width="2"/></svg>')
            edge = SceneNode('route', (160, 250, 100, 50), svg_draw(path, expected_sha256=digest(path.read_bytes())),
                             role='relation', relation=('a', 'b'))
            report = assemble_pptx(self.slide, ReconstructionScene(self.frame, [self.node('a'), self.node('b'), edge]))
            stream = io.BytesIO()
            self.deck.save(stream)
            reopened = Presentation(io.BytesIO(stream.getvalue())).slides[0]
            self.assertIn('cubicBezTo', reopened.shapes[-1]._element.xml)
            curve = reopened.shapes[-1]._element.xpath('.//a:cubicBezTo')[0]
            points = list(curve)
            self.assertAlmostEqual(int(points[0].get('x'))/int(points[2].get('x')), .9)
            self.assertAlmostEqual(int(points[1].get('x'))/int(points[2].get('x')), .1)
            move = reopened.shapes[-1]._element.xpath('.//a:moveTo/a:pt')[0]
            self.assertLess(int(points[0].get('y')), int(move.get('y')))
            self.assertEqual(points[2].get('y'), move.get('y'))
            self.assertEqual(report['relations'][0]['target'], 'b')
            self.assertEqual(report['visual_review'], 'NOT_PERFORMED')

    def test_text_scales_with_source_and_stays_transparent(self):
        node = SceneNode('label', (150, 250, 90, 30), text_draw('Protein', font_size_px=12), role='text')
        assemble_pptx(self.slide, ReconstructionScene(self.frame, [node]))
        text = self.slide.shapes[0]
        self.assertAlmostEqual(text.text_frame.paragraphs[0].font.size.pt, 12*.02*72, places=2)
        self.assertEqual(text.text, 'Protein')
        self.assertNotIn('<a:solidFill>', text._element.spPr.xml)

    def test_reused_text_callback_uses_new_scene_scale(self):
        draw = text_draw('Protein', font_size_px=12)
        small_frame = SourceFrame(self.frame.source, (1, 1, 4, 2))
        node = SceneNode('label', (150, 250, 90, 30), draw, role='text')
        assemble_pptx(self.slide, ReconstructionScene(small_frame, [node]))
        self.assertAlmostEqual(self.slide.shapes[0].text_frame.paragraphs[0].font.size.pt, 8.64, places=2)

    def test_protect_existing_slide(self):
        assemble_pptx(self.slide, ReconstructionScene(self.frame, [self.node('a')]))
        with self.assertRaisesRegex(ValueError, 'empty disposable'):
            assemble_pptx(self.slide, ReconstructionScene(self.frame, [self.node('b')]))
        self.assertEqual(len(self.slide.shapes), 1)

    def test_bad_source_bounds(self):
        for value in ((0, 0, 5, 5), (120, 220, 0, 3), (120, float('nan'), 1, 1)):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.frame.place(value)


if __name__ == '__main__':
    unittest.main()
