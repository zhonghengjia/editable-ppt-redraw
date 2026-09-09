import hashlib
import importlib.util
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from native_vectors import SVGProfileError, add_svg_component, add_rect_link, read_vectors, transform, multiply
from vendor.svg_paths.drawingml_paths import parse_svg_path, svg_path_to_absolute, normalize_path_commands
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

spec = importlib.util.spec_from_file_location('connection_audit', ROOT/'scripts/audit-pptx-connections.py')
auditor = importlib.util.module_from_spec(spec); spec.loader.exec_module(auditor)


class NativeToolkitTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.deck = Presentation(); self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def tearDown(self):
        self.temp.cleanup()

    def svg(self, body, attrs=''):
        source = self.folder/'source.svg'
        source.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" {attrs}>{body}</svg>', encoding='utf-8')
        return source

    def nodes(self):
        first = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        second = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(2), Inches(1))
        first.name, second.name = 'assessment', 'treatment'
        return first, second

    def save(self):
        target = self.folder/'check.pptx'; self.deck.save(target)
        Presentation(target)
        return target

    def test_all_bundled_icons_have_editable_curves_and_unique_ids(self):
        for index, source in enumerate(sorted((ROOT/'assets/lucide').glob('*.svg'))):
            result = add_svg_component(self.slide, source, index*2, 0, 1, 1, name=source.stem)
            self.assertGreater(result['paths'], 1)
        target = self.save()
        with ZipFile(target) as z:
            xml = z.read('ppt/slides/slide1.xml').decode()
            self.assertNotIn('<p:pic', xml)
            self.assertIn('<a:cubicBezTo>', xml)
            self.assertIn('cap="rnd"', xml)
            self.assertFalse(any('/media/' in name for name in z.namelist()))
        ids = [elem.get('id') for elem in self.slide._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')]
        self.assertEqual(len(ids), len(set(ids)))

    def test_closed_subpath_resets_quadratic_origin(self):
        result = normalize_path_commands(svg_path_to_absolute(parse_svg_path('M10 10 L20 20 Z Q13 10 16 10')))
        self.assertEqual(result[-1].args, [12, 10, 14, 10, 16, 10])

    def test_source_mapping_resolves_after_reopen(self):
        source = self.svg('<path id="tube" d="M1 1L5 5"/><circle id="end" cx="5" cy="5" r="2"/>')
        result = add_svg_component(self.slide, source, 1, 1, 2, 2, name='instrument')
        saved = self.save()
        reopened = Presentation(saved)
        group = next(s for s in reopened.slides[0].shapes if s.shape_id == result['group_id'])
        actual = {s.shape_id: s for s in group.shapes}
        self.assertEqual(len(result['element_map']), 2)
        for entry in result['element_map']:
            shape = actual[entry['shape_id']]
            self.assertEqual(shape.name, entry['output_name'])
            for measured, recorded in zip((shape.left, shape.top, shape.width, shape.height), entry['bounds_inches']):
                self.assertAlmostEqual(measured / 914400, recorded, places=5)

    def test_all_curve_commands_and_arc_degeneracies(self):
        result = normalize_path_commands(svg_path_to_absolute(parse_svg_path('m1 1 c1 0 1 1 2 1 s1 1 2 1 q1 1 2 2 t2 1 a2 3 20 0 1 2 3 z')))
        self.assertTrue(all(c.cmd in 'MLCZ' for c in result))
        zero = normalize_path_commands(svg_path_to_absolute(parse_svg_path('M0 0 A0 2 0 0 1 5 5')))
        self.assertEqual(zero[-1].cmd, 'L')

    def test_malformed_paths_fail(self):
        for d in ('M1', 'L1 2', 'M1 2 garbage', 'M1 2 L3', 'M1 2 Z3', 'M1e999 2', 'M0 0 A1 2 0 2 0 5 5'):
            with self.subTest(d=d), self.assertRaises(ValueError):
                parse_svg_path(d)

    def test_unsupported_content_fails_before_slide_mutation(self):
        for body in ('<image href="https://example.org/private.png"/>', '<text>editable?</text>',
                     '<path d="M1 1L2 2" filter="url(#x)"/>', '<use href="#x"/>', '<script/>',
                     '<rect width="4" height="4" style="opacity:0.2"/>', '<path d="M0 0L2 2" fill-rule="evenodd"/>'):
            before = len(self.slide.shapes)
            with self.subTest(body=body), self.assertRaises(SVGProfileError):
                add_svg_component(self.slide, self.svg(body), 0, 0, 1, 1)
            self.assertEqual(len(self.slide.shapes), before)

    def test_entities_duplicates_and_anisotropy_rejected(self):
        inputs = ['<!DOCTYPE svg [<!ENTITY s "x">]><svg/>',
                  '<svg viewBox="0 0 24 24"><path id="a" d="M1 1L2 2"/><path id="a" d="M2 2L3 3"/></svg>',
                  '<svg viewBox="0 0 24 24"><g transform="scale(2 3)"><path d="M1 1L2 2"/></g></svg>']
        for text in inputs:
            source = self.folder/'invalid.svg'; source.write_text(text)
            with self.assertRaises(SVGProfileError):
                read_vectors(source)

    def test_nested_transform_composition_and_style_inheritance(self):
        source = self.svg('<g transform="scale(2)"><g transform="translate(3 4)" stroke="red"><path d="M1 1L2 2"/></g></g>', 'fill="none"')
        specs, _, _ = read_vectors(source)
        self.assertEqual(specs[0]['commands'][0].args, [8, 10])
        self.assertIsNone(specs[0]['fill']); self.assertEqual(specs[0]['stroke'], 'FF0000')
        self.assertAlmostEqual(transform('rotate(90)')[1], 1)

    def test_contain_placement_and_nonzero_viewbox(self):
        source = self.folder/'offset.svg'
        source.write_text('<svg viewBox="10 20 20 10"><rect x="10" y="20" width="20" height="10"/></svg>')
        group = add_svg_component(self.slide, source, 1, 1, 2, 2)['group']
        self.assertEqual(group.left, Inches(1)); self.assertEqual(group.top, Inches(1.5))
        self.assertEqual(group.width, Inches(2)); self.assertEqual(group.height, Inches(1))

    def test_bound_link_actual_export_and_semantic_arrow_check(self):
        first, second = self.nodes()
        add_rect_link(self.slide, first, second, name='flow-main')
        expected = [{'name':'flow-main', 'source':'assessment', 'target':'treatment', 'end_arrow':'triangle', 'start_arrow':'none'}]
        report = auditor.audit(self.save(), True, expected)
        self.assertTrue(report['ok'], report); self.assertTrue(report['geometry_complete'])
        expected[0]['source'] = 'treatment'
        self.assertFalse(auditor.audit(self.save(), True, expected)['ok'])

    def test_reverse_upward_connection_preserves_end_arrow(self):
        first, second = self.nodes()
        add_rect_link(self.slide, second, first, source_port='top', target_port='bottom')
        report = auditor.audit(self.save(), True)
        self.assertTrue(report['ok'], report)
        self.assertEqual(report['connectors'][0]['end_arrow'], 'triangle')

    def test_unaligned_or_inward_links_fail_before_mutation(self):
        first, second = self.nodes()
        before = len(self.slide.shapes)
        with self.assertRaises(ValueError):
            add_rect_link(self.slide, first, second, source_port='top', target_port='bottom')
        second.left += Inches(.1)
        with self.assertRaises(ValueError):
            add_rect_link(self.slide, first, second)
        self.assertEqual(before, len(self.slide.shapes))

    def test_broken_endpoint_and_off_port_detected(self):
        first, second = self.nodes(); link = add_rect_link(self.slide, first, second)
        link._element.spPr.xfrm.off.set('x', str(Inches(3)))
        report = auditor.audit(self.save(), True)
        self.assertIn('endpoint_off_port', [i['code'] for i in report['issues']])
        endpoint = link._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}endCxn')
        endpoint.set('id', '99999')
        self.assertIn('dangling_endpoint', [i['code'] for i in auditor.audit(self.save(), True)['issues']])

    def test_unbound_missing_and_empty_connector_cases(self):
        self.assertFalse(auditor.audit(self.save(), True)['ok'])
        first, second = self.nodes(); link = add_rect_link(self.slide, first, second)
        endpoint = link._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}endCxn')
        endpoint.getparent().remove(endpoint)
        report = auditor.audit(self.save(), True, [{'name': 'missing'}])
        codes = [i['code'] for i in report['issues']]
        self.assertIn('unbound_endpoint', codes); self.assertIn('missing_expected_connector', codes)
        self.assertFalse(report['geometry_complete'])

    def test_actual_diagonal_and_zero_length_connectors_are_detected(self):
        first, second = self.nodes(); link = add_rect_link(self.slide, first, second)
        ext = link._element.spPr.xfrm.ext
        ext.set('cx', str(Inches(.1)))
        self.assertIn('diagonal_connector', [i['code'] for i in auditor.audit(self.save(), True)['issues']])
        ext.set('cx', '0'); ext.set('cy', '0')
        self.assertIn('zero_length_connector', [i['code'] for i in auditor.audit(self.save(), True)['issues']])

    def test_rotated_geometry_is_explicitly_unverified(self):
        first, second = self.nodes(); link = add_rect_link(self.slide, first, second)
        link._element.spPr.xfrm.set('rot', '900000')
        report = auditor.audit(self.save(), True)
        self.assertTrue(report['ok'])
        self.assertFalse(report['geometry_complete'])
        self.assertEqual(len(report['geometry_unverified']), 1)

    def test_cli_appends_to_template_without_overwriting(self):
        self.slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'KEEP ORIGINAL 中文'
        template = self.save(); digest = hashlib.sha256(template.read_bytes()).hexdigest()
        output = self.folder/'new.pptx'
        command = [sys.executable, '-B', str(ROOT/'scripts/native_vectors.py'), str(ROOT/'assets/lucide/stethoscope.svg'), str(output), '--template', str(template)]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        reopened = Presentation(output)
        self.assertEqual(len(reopened.slides), 2)
        self.assertIn('KEEP ORIGINAL 中文', reopened.slides[0].shapes[0].text)
        self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(), digest)
        output_hash = hashlib.sha256(output.read_bytes()).hexdigest()
        self.assertNotEqual(subprocess.run(command, capture_output=True).returncode, 0)
        self.assertEqual(hashlib.sha256(output.read_bytes()).hexdigest(), output_hash)


if __name__ == '__main__':
    unittest.main()
