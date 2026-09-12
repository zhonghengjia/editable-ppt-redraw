"""Production dependency tests, using actual extraction, tracing and PPTX readback."""
import copy
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image, ImageDraw
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import reconstruction_pipeline as P

NODE = os.environ.get('EDITABLE_PPT_NODE') or shutil.which('node')


def example_manifest(source):
    return dict(schema_version=1, mode='faithful', execution_profile='standard',
        source=dict(path=source.name, width=180, height=120, sha256=P.sha(source)),
        canvas=dict(width=180, height=120), targets=[dict(format='pptx', path='candidate.pptx')],
        modules=[dict(id='panel', type='illustration', bbox=[0, 0, 180, 120], reading_order=1)],
        connections=[], uncertainties=[], construction=dict(version=1, slide_width_inches=6),
        source_inventory=[
            dict(id='body', module='panel', role='shape', bbox=[10, 15, 145, 88],
                representation='native_composite', construction=dict(kind='trace',
                    observation='Synthetic two-tone body, white hole and thin projecting tip',
                    parent='background', extract=dict(method='selector', foreground=[[40, 40]],
                        background=[], reserved_regions=[], selector=dict(rgb=[[0, 100], [60, 150], [100, 210]],
                            differences=[], observation='Synthetic blue tones, exclude white')),
                    trace=dict(colors=4, representation='palette_stack', max_native_paths=30, max_native_commands=2000))),
            dict(id='label', module='panel', role='text', bbox=[20, 102, 110, 16], text='Observed body',
                representation='native_primitive', construction=dict(kind='text', observation='Source label below body',
                    options=dict(font_size_px=12), parent='background')),
            dict(id='background', module='panel', role='shape', bbox=[0, 0, 180, 120],
                representation='native_primitive', construction=dict(kind='rectangle', container=True,
                    observation='White source panel', options=dict(fill='FFFFFF', stroke=None)))])


class ProductionRoute(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.png'
        image = Image.new('RGB', (180, 120), 'white')
        d = ImageDraw.Draw(image)
        d.ellipse((25, 25, 120, 90), fill=(30, 100, 170))
        d.ellipse((35, 35, 90, 70), fill=(60, 130, 190))
        d.ellipse((55, 50, 66, 61), fill='white')
        d.line((106, 37, 142, 18), fill=(30, 100, 170), width=2)
        image.save(self.source)
        self.manifest = example_manifest(self.source)
        self.path = self.root / 'manifest.json'
        self.path.write_text(json.dumps(self.manifest), encoding='utf-8')

    def prepare(self, suffix='prepared'):
        self.path.write_text(json.dumps(self.manifest), encoding='utf-8')
        out = self.root / suffix
        P.prepare(self.path, out, node=NODE)
        return out

    def review(self, prepared):
        # This asserts synthetic fixture support, not fabricated real-image review.
        review = P.read_json(prepared / 'review-template.json')
        for item in review['items']:
            item.update(accepted=True, note='Synthetic test fixture: expected two-tone body with hole and tip')
        path = self.root / (prepared.name + '-review.json')
        path.write_text(json.dumps(review), encoding='utf-8')
        return path

    def test_actual_source_extraction_trace_and_grouped_pptx(self):
        before = P.sha(self.source)
        prepared = self.prepare()
        output = self.root / 'build'
        report = P.build(prepared, output, review=self.review(prepared), node=NODE)
        deck = Presentation(output / 'candidate.pptx')
        shapes = {s.name: s for s in deck.slides[0].shapes}
        self.assertEqual(list(shapes), ['background', 'body', 'label'])
        self.assertGreater(len(shapes['body'].shapes), 1)
        self.assertEqual(shapes['label'].text, 'Observed body')
        self.assertAlmostEqual(shapes['label'].text_frame.paragraphs[0].font.size.pt, 28.8, places=2)
        self.assertFalse(deck.slides[0]._element.xpath('.//p:pic'))
        colors = set(shapes['body']._element.xpath('.//a:srgbClr/@val'))
        self.assertGreaterEqual(len(colors), 2)
        self.assertNotEqual(report['items'][1]['source_bbox'], self.manifest['source_inventory'][0]['bbox'])
        source_bbox = report['items'][1]['source_bbox']
        self.assertGreater(source_bbox[0], 10)
        self.assertGreaterEqual(source_bbox[0] + source_bbox[2], 143)
        resolved = P.read_json(output / 'resolved-manifest.json')
        for original, item in zip(self.manifest['source_inventory'], resolved['source_inventory']):
            self.assertEqual(original['id'], item['id'])
            self.assertEqual(original['bbox'], item['bbox'])
            self.assertEqual(str(shapes[item['id']].shape_id), item['output_id'])
        self.assertEqual(P.sha(self.source), before)
        self.assertIsNone(P.read_json(output / 'quality-report.json')['delivery_ready'])
        self.assertEqual(report['visual_review'], 'NOT_PERFORMED')

    def test_all_inventory_items_must_have_recipe(self):
        del self.manifest['source_inventory'][0]['construction']
        self.assertTrue(P.validate_contract(self.manifest))
        self.path.write_text(json.dumps(self.manifest))
        with self.assertRaisesRegex(ValueError, 'Invalid manifest'):
            P.prepare(self.path, self.root / 'bad', node=NODE)
        self.assertFalse((self.root / 'bad').exists())

    def test_source_hash_and_frame_not_inferred_from_output(self):
        for change in ('hash', 'frame'):
            with self.subTest(change=change):
                manifest = copy.deepcopy(self.manifest)
                if change == 'hash':
                    manifest['source']['sha256'] = '0'*64
                else:
                    manifest['canvas']['width'] = 190
                self.path.write_text(json.dumps(manifest))
                with self.assertRaises(ValueError):
                    P.prepare(self.path, self.root / change, node=NODE)
                self.assertFalse((self.root / change).exists())

    def test_no_build_without_extraction_review(self):
        prepared = self.prepare()
        with self.assertRaisesRegex(ValueError, 'inspect extracted'):
            P.build(prepared, self.root / 'bad', node=NODE)
        self.assertFalse((self.root / 'bad').exists())

    def test_rejected_and_stale_reviews_are_distinct(self):
        prepared = self.prepare()
        review = self.review(prepared)
        data = P.read_json(review)
        data['items'][0]['accepted'] = False
        review.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'inspection rejected'):
            P.build(prepared, self.root / 'bad', review=review, node=NODE)
        data['items'][0]['candidate_sha256'] = '0'*64
        review.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'stale component'):
            P.build(prepared, self.root / 'bad', review=review, node=NODE)
        self.assertFalse((self.root / 'bad').exists())

    def native_manifest(self):
        from appearance_fidelity import sample_paint
        item = self.manifest['source_inventory'][0]
        item['bbox'] = [25, 25, 95, 65]
        item['construction'] = dict(kind='native_component', parent='background',
            observation='Synthetic source-sampled gradient geometry')
        fill = sample_paint(self.source, dict(source_sha256=P.sha(self.source), kind='linear', angle=0,
            positions=[0, 1], patches=[[45, 44, 4, 4], [102, 50, 4, 4]]))
        self.manifest['native_components'] = [dict(id='body', output_name='body', viewbox=[0, 0, 95, 65],
            parts=[dict(id='surface', role='surface', observation='Synthetic full-bound surface',
                d='M0 0 L95 0 L95 65 L0 65 Z', fill=fill)])]

    def test_native_paint_pipeline_reuses_frozen_source_and_preserves_gradient(self):
        self.native_manifest()
        prepared = self.prepare('paint')
        # Must not reread this unrelated/current source once preparation is frozen.
        self.source.write_bytes(b'changed after preparation')
        report = P.build(prepared, self.root / 'paint-build')
        slide = Presentation(self.root / 'paint-build/candidate.pptx').slides[0]
        group = slide.shapes[1]
        self.assertEqual(group.name, 'body')
        self.assertEqual(len(group._element.xpath('.//a:gradFill')), 1)
        self.assertEqual(group._element.xpath('.//a:gs/@pos'), ['0', '100000'])
        self.assertEqual(group._element.xpath('.//a:gs/a:srgbClr/@val'), ['3C82BE', '1E64AA'])
        self.assertFalse(slide._element.xpath('.//p:pic'))
        self.assertEqual(report['items'][1]['source_bbox'], [25, 25, 95, 65])
        self.assertEqual(report['visual_review'], 'NOT_PERFORMED')

    def test_native_sample_tampering_fails_without_partial_pptx(self):
        self.native_manifest()
        self.manifest['native_components'][0]['parts'][0]['fill']['stops'][0]['color'] = '000000'
        prepared = self.prepare('paint-tampered')
        with self.assertRaisesRegex(ValueError, 'differs'):
            P.build(prepared, self.root / 'failed-paint')
        self.assertFalse((self.root / 'failed-paint/candidate.pptx').exists())

    def test_native_component_has_one_authority_and_cannot_be_ignored(self):
        self.native_manifest()
        self.assertEqual(P.validate_contract(self.manifest), [])
        bad = copy.deepcopy(self.manifest)
        bad['native_components'][0]['output_name'] = 'other'
        self.assertTrue(P.validate_contract(bad))
        bad = copy.deepcopy(self.manifest)
        bad['source_inventory'][0]['construction'] = dict(kind='rectangle', options={}, observation='Wrong route')
        self.assertTrue(P.validate_contract(bad))
        bad = copy.deepcopy(self.manifest)
        bad['native_components'] = []
        self.assertTrue(P.validate_contract(bad))

    def test_annotations_are_only_part_of_grabcut_recipe(self):
        extract = self.manifest['source_inventory'][0]['construction']['extract']
        extract['annotations'] = []
        self.assertTrue(P.validate_contract(self.manifest))
        extract['method'] = 'grabcut'
        extract['background'] = [[5, 5]]
        del extract['selector']
        self.assertEqual(P.validate_contract(self.manifest), [])

    def test_tampered_alpha_manifest_and_svg_are_rejected(self):
        for name in ('alpha', 'manifest'):
            with self.subTest(name=name):
                prepared = self.prepare(name)
                review = self.review(prepared)
                target = prepared / ('item-0000/alpha.png' if name == 'alpha' else 'input-manifest.json')
                target.write_bytes(target.read_bytes() + b'changed')
                with self.assertRaisesRegex(ValueError, 'changed'):
                    P.build(prepared, self.root / (name + '-build'), review=review, node=NODE)
        # Real SVG import with a source-observed nonzero viewBox and a curve.
        vector = self.root / 'observed.svg'
        vector.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="10 20 100 50"><path id="body" d="M10 45 C20 20 100 70 110 45" fill="none" stroke="#123456"/></svg>')
        self.manifest['source_inventory'][0]['construction'] = dict(kind='svg', observation='Synthetic observed curve',
            asset=dict(path=vector.name, sha256=P.sha(vector)))
        prepared = self.prepare('svg')
        (prepared / 'item-0000/source.svg').write_text('changed')
        with self.assertRaisesRegex(ValueError, 'hash changed'):
            P.build(prepared, self.root / 'svg-build', node=NODE)

    def test_correction_rebuild_does_not_append_shapes(self):
        first = self.prepare('first')
        P.build(first, self.root / 'b1', review=self.review(first), node=NODE)
        self.manifest['source_inventory'][1]['text'] = 'Revised text'
        second = self.prepare('second')
        with self.assertRaisesRegex(ValueError, 'different preparation'):
            P.build(second, self.root / 'stale', review=self.review(first), node=NODE)
        P.build(second, self.root / 'b2', review=self.review(second), node=NODE)
        old = Presentation(self.root / 'b1/candidate.pptx').slides[0]
        new = Presentation(self.root / 'b2/candidate.pptx').slides[0]
        self.assertEqual(len(old.shapes), len(new.shapes))
        self.assertEqual(old.shapes[-1].text, 'Observed body')
        self.assertEqual(new.shapes[-1].text, 'Revised text')
        with self.assertRaisesRegex(ValueError, 'new build'):
            P.build(second, self.root / 'b2', review=self.review(second), node=NODE)

    def test_native_budget_failure_never_saves_partial_pptx(self):
        self.manifest['source_inventory'][0]['construction']['trace']['max_native_commands'] = 1
        prepared = self.prepare()
        output = self.root / 'failed'
        with self.assertRaisesRegex(ValueError, 'budget'):
            P.build(prepared, output, review=self.review(prepared), node=NODE)
        self.assertFalse((output / 'candidate.pptx').exists())
        self.assertEqual(P.read_json(output / 'failure.json')['status'], 'FAILED')
        self.assertTrue(list(output.glob('*.trace.failed.json')))

    def test_missing_node_is_explicit_and_no_output(self):
        with patch.object(P.shutil, 'which', return_value=None):
            with self.assertRaisesRegex(RuntimeError, 'no automatic'):
                P.prepare(self.path, self.root / 'bad')
        self.assertFalse((self.root / 'bad').exists())

    def test_text_and_rectangle_do_not_require_segmentation(self):
        self.manifest['source_inventory'].pop(0)
        prepared = self.prepare()
        P.build(prepared, self.root / 'simple', node=NODE)
        slide = Presentation(self.root / 'simple/candidate.pptx').slides[0]
        self.assertEqual([s.name for s in slide.shapes], ['background', 'label'])

    def test_relation_items_have_real_endpoints_and_module_mapping(self):
        edge = dict(id='edge', module='panel', role='connector', representation='native_composite',
            bbox=[10, 10, 20, 20], construction=dict(kind='svg', observation='Measured relationship',
                asset=dict(path='edge.svg', sha256='1'*64), connection='link', endpoints=['body', 'background']))
        self.manifest['source_inventory'].append(edge)
        self.manifest['connections'] = [dict(id='link', source='panel', target='panel')]
        self.assertEqual(P.validate_contract(self.manifest), [])
        edge['construction']['endpoints'][1] = 'missing'
        self.assertTrue(P.validate_contract(self.manifest))
        edge['construction']['endpoints'][1] = 'background'
        self.manifest['connections'][0]['target'] = 'wrong-module'
        self.assertTrue(P.validate_contract(self.manifest))

    def test_wrong_route_and_unknown_fields_do_not_fallback(self):
        for change in (dict(kind='template'), dict(kind='rectangle'), dict(options={}), dict(container=False)):
            data = copy.deepcopy(self.manifest)
            data['source_inventory'][0]['construction'].update(change)
            self.assertTrue(P.validate_contract(data))

    def test_isolated_extraction_process_uses_existing_protocol(self):
        P.prepare(self.path, self.root / 'isolated', node=NODE, extract_python=sys.executable)
        P.build(self.root / 'isolated', self.root / 'isolated-build',
            review=self.review(self.root / 'isolated'), node=NODE)
        self.assertTrue((self.root / 'isolated-build/candidate.pptx').is_file())

    def test_vector_curve_uses_same_frame_without_segmentation(self):
        vector = self.root / 'observed.svg'
        vector.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="10 20 100 50"><path id="fold" d="M10 45 C20 20 100 70 110 45" fill="none" stroke="#123456"/></svg>')
        item = self.manifest['source_inventory'][0]
        item['bbox'] = [40, 40, 100, 50]
        item['construction'] = dict(kind='svg', observation='Synthetic observed nonzero-origin curve',
            asset=dict(path=vector.name, sha256=P.sha(vector)), parent='background')
        self.manifest['targets'][0]['path'] = 'Named result.pptx'
        prepared = self.prepare('vector')
        report = P.build(prepared, self.root / 'vector-build', node=NODE)
        slide = Presentation(self.root / 'vector-build/Named result.pptx').slides[0]
        group = slide.shapes[1]
        self.assertEqual(group.name, 'body')
        self.assertTrue(group._element.xpath('.//a:cubicBezTo'))
        self.assertAlmostEqual(group.left / 914400, 40 / 30, places=5)
        self.assertEqual(report['items'][1]['source_bbox'], [40, 40, 100, 50])

    def test_preparation_failure_leaves_no_success_receipt(self):
        self.manifest['source_inventory'][0]['construction']['extract']['foreground'] = [[0, 0]]
        self.path.write_text(json.dumps(self.manifest))
        with self.assertRaises(ValueError):
            P.prepare(self.path, self.root / 'failed-prep', node=NODE)
        self.assertFalse((self.root / 'failed-prep/prepared.json').exists())
        self.assertEqual(P.read_json(self.root / 'failed-prep/failure.json')['phase'], 'prepare')

    def test_supplied_bad_svg_is_not_silently_replaced(self):
        vector = self.root / 'unsupported.svg'
        vector.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><text x="0" y="8">Not native text</text></svg>')
        self.manifest['source_inventory'][0]['construction'] = dict(kind='svg', observation='Unsupported outlined source label',
            asset=dict(path=vector.name, sha256=P.sha(vector)))
        self.path.write_text(json.dumps(self.manifest))
        with self.assertRaises(ValueError):
            P.prepare(self.path, self.root / 'unsupported', node=NODE)
        self.assertFalse((self.root / 'unsupported/prepared.json').exists())

    def test_roundtrip_preserves_source_and_semantic_manifest_sections(self):
        self.manifest['semantic_constraints'] = [dict(id='label-preserved', subjects=['label'],
            type='must_preserve_text', rule='Keep exact observed label', verification='Native text readback')]
        prepared = self.prepare()
        P.build(prepared, self.root / 'roundtrip', review=self.review(prepared), node=NODE)
        resolved = P.read_json(self.root / 'roundtrip/resolved-manifest.json')
        for field in ('modules', 'connections', 'semantic_constraints', 'uncertainties'):
            self.assertEqual(resolved[field], self.manifest[field])


if __name__ == '__main__':
    unittest.main()
