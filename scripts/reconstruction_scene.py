"""Source-frame assembly over existing emitters; no alternate SVG/Paint engine.

One scene owns placement and paint order, while the manifest owns observations,
component content and authorization. Callbacks allow a qualified target adapter
without duplicating the native-component schema. Failed assembly is discarded,
never delivered as a partial success.
"""
from __future__ import annotations

from dataclasses import dataclass, field
import heapq
import math
from pathlib import Path
from typing import Callable


def _rect(value):
    if (len(value) != 4 or any(isinstance(v, bool) or not isinstance(v, (int, float))
                              or not math.isfinite(v) for v in value)
            or min(value[2:]) <= 0):
        raise ValueError('rectangle requires finite [x,y,width,height] with positive size')
    return tuple(value)


@dataclass(frozen=True)
class SourceFrame:
    """Source region in pixels -> target viewport in inches, uniformly contained."""
    source: tuple
    viewport: tuple

    @property
    def scale(self):
        _, _, sw, sh = _rect(self.source)
        _, _, tw, th = _rect(self.viewport)
        return min(tw/sw, th/sh)

    def place(self, bbox):
        sx, sy, sw, sh = _rect(self.source)
        tx, ty, tw, th = _rect(self.viewport)
        x, y, w, h = _rect(bbox)
        if x < sx or y < sy or x+w > sx+sw+1e-8 or y+h > sy+sh+1e-8:
            raise ValueError('object bbox outside source frame')
        scale = self.scale
        return (tx+(tw-sw*scale)/2+(x-sx)*scale,
                ty+(th-sh*scale)/2+(y-sy)*scale, w*scale, h*scale)


@dataclass
class SceneNode:
    id: str
    bbox: tuple
    draw: Callable
    role: str = 'component'
    parent: str | None = None
    behind: tuple = ()
    z: int = 0
    relation: tuple | None = None


@dataclass
class ReconstructionScene:
    frame: SourceFrame
    nodes: list[SceneNode] = field(default_factory=list)

    def ordered(self):
        """Stable topological order; explicit dependencies take priority over z.

        A parent is a container background, not arbitrary anatomical grouping.
        Interleaved anatomy uses explicit behind edges and existing part trees.
        """
        nodes = self.nodes
        if not nodes or len(nodes) > 10000:
            raise ValueError('scene requires 1..10000 nodes')
        by = {}
        for node in nodes:
            if not isinstance(node.id, str) or not node.id.strip() or node.id in by:
                raise ValueError('scene IDs must be nonempty and unique')
            if not callable(node.draw) or type(node.z) is not int:
                raise ValueError('invalid draw callback or z index')
            if node.role not in ('component', 'text', 'container', 'relation'):
                raise ValueError('unknown scene role')
            self.frame.place(node.bbox)
            by[node.id] = node
        edges = {n.id: set() for n in nodes}
        for node in nodes:
            if node.parent is not None:
                if node.parent not in by or by[node.parent].role != 'container':
                    raise ValueError('parent must reference a container background')
                px, py, pw, ph = by[node.parent].bbox
                x, y, w, h = node.bbox
                if x < px or y < py or x+w > px+pw+1e-8 or y+h > py+ph+1e-8:
                    raise ValueError('child geometry outside declared container')
                edges[node.parent].add(node.id)
            if not isinstance(node.behind, (list, tuple)):
                raise ValueError('behind requires a list of scene IDs')
            for target in node.behind:
                if target not in by:
                    raise ValueError('unknown layer target: ' + str(target))
                edges[node.id].add(target)
            if node.role == 'relation':
                if (not isinstance(node.relation, (tuple, list)) or len(node.relation) != 2
                        or any(v not in by or by[v].role == 'relation' for v in node.relation)):
                    raise ValueError('relation requires existing non-relation source/target IDs')
            elif node.relation is not None:
                raise ValueError('relation endpoints require relation role')
        indegree = {n.id: 0 for n in nodes}
        for targets in edges.values():
            for target in targets:
                indegree[target] += 1
        rank = {n.id: i for i, n in enumerate(nodes)}
        ready = [(n.z, rank[n.id], n.id) for n in nodes if not indegree[n.id]]
        heapq.heapify(ready)
        result = []
        while ready:
            _, _, identity = heapq.heappop(ready)
            result.append(by[identity])
            for target in sorted(edges[identity], key=rank.get):
                indegree[target] -= 1
                if not indegree[target]:
                    heapq.heappush(ready, (by[target].z, rank[target], target))
        if len(result) != len(nodes):
            raise ValueError('cyclic layer/containment constraints; do not guess order')
        return result


def assemble_pptx(slide, scene):
    """Construct on an empty disposable slide; return actual top-level mappings.

    Do not save/deliver on exception. Callbacks return one top-level shape/group;
    structural output is inspected here but visual semantics remain source review.
    """
    if len(slide.shapes):
        raise ValueError('assemble on an empty disposable slide, preserving existing artifacts')
    ordered = scene.ordered()  # all placement/dependency errors before mutation
    inventory, relations = [], []
    for node in ordered:
        before = len(slide.shapes)
        bounds = scene.frame.place(node.bbox)
        shape = node.draw(slide, bounds, node.id, scene.frame.scale)
        if (len(slide.shapes) != before+1
                or getattr(shape, 'shape_id', None) != slide.shapes[-1].shape_id):
            raise ValueError('adapter must emit exactly one top-level shape/group; discard failed slide')
        shape.name = node.id
        inventory.append({'id': node.id, 'role': node.role, 'source_bbox': list(node.bbox),
                          'shape_id': shape.shape_id, 'output_name': shape.name,
                          'placed_bounds_inches': list(bounds), 'parent_background': node.parent})
        if node.relation is not None:
            relations.append({'id': node.id, 'source': node.relation[0], 'target': node.relation[1],
                              'shape_id': shape.shape_id, 'binding': 'source-observed path, not drag-bound'})
    names = [el.get('name') for el in slide._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')]
    if len(names) != len(set(names)):
        raise ValueError('duplicate output names; discard failed slide')
    return {'items': inventory, 'relations': relations, 'paint_order': [n.id for n in ordered],
            'source_frame': list(scene.frame.source), 'viewport_inches': list(scene.frame.viewport),
            'visual_review': 'NOT_PERFORMED'}


def svg_draw(path, *, expected_sha256):
    """Keep source controls/viewBox through the existing native SVG emitter."""
    from native_vectors import add_svg_component
    from raster_components import digest
    path = Path(path)

    def draw(slide, bbox, identity, source_scale):
        if digest(path.read_bytes()) != expected_sha256:
            raise ValueError('SVG source hash changed')
        return add_svg_component(slide, path, *bbox, name=identity)['group']
    return draw


def manifest_component_draw(manifest, component_id, *, base_dir):
    from native_components import add_manifest_component

    def draw(slide, bbox, identity, source_scale):
        matches = [c for c in manifest.get('native_components', []) if c['id'] == component_id]
        if len(matches) != 1 or matches[0]['output_name'] != identity:
            raise ValueError('scene identity must equal manifest component output_name')
        return add_manifest_component(slide, manifest, component_id, *bbox, base_dir=base_dir)['group']
    return draw


def text_draw(text, *, font_size_px, color='222222', font='Arial', bold=False, align='left'):
    """Text sizes use the same source-to-target scale as geometry, not fixed pt."""
    if (not isinstance(text, str) or not text or not math.isfinite(font_size_px)
            or font_size_px <= 0 or align not in ('left', 'center', 'right')):
        raise ValueError('invalid native text specification')
    def draw(slide, bbox, identity, source_scale):
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        shape = slide.shapes.add_textbox(*(Inches(v) for v in bbox))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split('\n')):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = line
            para.space_before = para.space_after = Pt(0)
            para.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
            para.font.name, para.font.size, para.font.bold = font, Pt(font_size_px*source_scale*72), bold
            para.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
        return shape
    return draw


def rectangle_draw(*, fill=None, stroke='222222', stroke_pt=1):
    def draw(slide, bbox, identity, source_scale):
        from pptx.util import Inches, Pt
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(v) for v in bbox))
        for target, value in ((shape.fill, fill), (shape.line.fill, stroke)):
            if value is None:
                target.background()
            else:
                target.solid()
                target.fore_color.rgb = RGBColor.from_string(value.lstrip('#'))
        shape.line.width = Pt(stroke_pt)
        return shape
    return draw
