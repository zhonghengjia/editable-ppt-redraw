"""Local SVG vector components and bound PowerPoint links, not a pixel recognizer.

Only the explicit SVG profile below is accepted. Preflight is complete before
the destination slide is mutated. No networking, raster fallback or model calls.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.sax.saxutils import quoteattr

from PIL import ImageColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt
from native_paint import paint_xml

from vendor.svg_paths.drawingml_paths import (
    PathCommand, normalize_path_commands, parse_svg_path,
    path_commands_to_drawingml, svg_path_to_absolute,
)

SVG_NS = 'http://www.w3.org/2000/svg'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
EMU = 914400
MAX_SVG_BYTES = 2_000_000
MAX_SVG_ELEMENTS = 2000
MAX_PATH_CHARS = 100_000
IDENTITY = (1., 0., 0., 1., 0., 0.)
STYLE = {'fill', 'stroke', 'stroke-width', 'stroke-linecap', 'stroke-linejoin',
         'color', 'fill-rule', 'fill-opacity', 'stroke-opacity'}
ATTRS = {
    'svg': {'width', 'height', 'viewBox', 'version'}, 'g': set(),
    'path': {'d'}, 'rect': {'x', 'y', 'width', 'height', 'rx', 'ry'},
    'circle': {'cx', 'cy', 'r'}, 'ellipse': {'cx', 'cy', 'rx', 'ry'},
    'line': {'x1', 'y1', 'x2', 'y2'},
    'polyline': {'points'}, 'polygon': {'points'},
}
NUMBER = r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?'
PORTS = {'top': 0, 'left': 1, 'bottom': 2, 'right': 3}


class SVGProfileError(ValueError):
    """Source needs explicit vector normalization, not silent conversion loss."""


def number(value, default=0.):
    value = str(default if value is None else value).strip()
    if not re.fullmatch(NUMBER + r'(?:px)?', value):
        raise SVGProfileError(f'Expected a numeric pixel length: {value[:80]}')
    result = float(value.removesuffix('px'))
    if not math.isfinite(result) or abs(result) > 1e9:
        raise SVGProfileError('Coordinate is nonfinite or exceeds the profile limit')
    return result


def numbers(value):
    tokens = re.findall(NUMBER, value)
    if re.sub(NUMBER, '', value).strip(' ,\t\r\n'):
        raise SVGProfileError('Invalid coordinate list')
    return [number(token) for token in tokens]


def multiply(m, n):
    a, b, c, d, e, f = m
    u, v, w, x, y, z = n
    return (a*u+c*v, b*u+d*v, a*w+c*x, b*w+d*x, a*y+c*z+e, b*y+d*z+f)


def transform(value):
    result, pos = IDENTITY, 0
    for match in re.finditer(r'(translate|scale|rotate)\s*\(([^)]*)\)', value):
        if value[pos:match.start()].strip(' ,\t\r\n'):
            raise SVGProfileError('Only translate, positive uniform scale and rotate are supported')
        kind, vals = match.group(1), numbers(match.group(2))
        if kind == 'translate' and len(vals) in (1, 2):
            local = (1, 0, 0, 1, vals[0], vals[1] if len(vals) == 2 else 0)
        elif kind == 'scale' and len(vals) in (1, 2):
            if vals[0] <= 0 or (len(vals) == 2 and not math.isclose(*vals)):
                raise SVGProfileError('Nonuniform or mirrored scaling must be normalized first')
            local = (vals[0], 0, 0, vals[0], 0, 0)
        elif kind == 'rotate' and len(vals) in (1, 3):
            angle = math.radians(vals[0]); c, s = math.cos(angle), math.sin(angle)
            local = (c, s, -s, c, 0, 0)
            if len(vals) == 3:
                x, y = vals[1:]
                local = multiply((1, 0, 0, 1, x, y), multiply(local, (1, 0, 0, 1, -x, -y)))
        else:
            raise SVGProfileError(f'Invalid {kind} arguments')
        result = multiply(result, local); pos = match.end()
    if value[pos:].strip(' ,\t\r\n'):
        raise SVGProfileError('Unsupported transform')
    return result


def paint(value, current_color):
    value = current_color if value == 'currentColor' else value
    if value == 'none':
        return None
    try:
        return '%02X%02X%02X' % ImageColor.getrgb(value)
    except (ValueError, TypeError):
        raise SVGProfileError(f'Unsupported solid color: {str(value)[:80]}') from None


def path_for(tag, elem):
    a = elem.attrib
    n = lambda key, default=0: number(a.get(key), default)
    if tag == 'path':
        return a.get('d', '')
    if tag == 'line':
        return f'M {n("x1")} {n("y1")} L {n("x2")} {n("y2")}'
    if tag in ('polygon', 'polyline'):
        pts = numbers(a.get('points', ''))
        if len(pts) < (6 if tag == 'polygon' else 4) or len(pts) % 2:
            raise SVGProfileError('Insufficient or incomplete polygon/polyline points')
        return 'M ' + ' '.join(map(str, pts)) + (' Z' if tag == 'polygon' else '')
    if tag in ('circle', 'ellipse'):
        x, y = n('cx'), n('cy')
        rx, ry = (n('r'), n('r')) if tag == 'circle' else (n('rx'), n('ry'))
        if min(rx, ry) <= 0:
            raise SVGProfileError('Circle/ellipse radii must be positive')
        return f'M {x-rx} {y} A {rx} {ry} 0 1 0 {x+rx} {y} A {rx} {ry} 0 1 0 {x-rx} {y} Z'
    x, y, w, h = n('x'), n('y'), n('width'), n('height')
    if min(w, h) <= 0:
        raise SVGProfileError('Rectangle dimensions must be positive')
    rx = n('rx', n('ry')); ry = n('ry', rx)
    if min(rx, ry) < 0:
        raise SVGProfileError('Rounded corner radii cannot be negative')
    rx, ry = min(rx, w/2), min(ry, h/2)
    if not rx or not ry:
        return f'M {x} {y} H {x+w} V {y+h} H {x} Z'
    return (f'M {x+rx} {y} H {x+w-rx} A {rx} {ry} 0 0 1 {x+w} {y+ry} '
            f'V {y+h-ry} A {rx} {ry} 0 0 1 {x+w-rx} {y+h} H {x+rx} '
            f'A {rx} {ry} 0 0 1 {x} {y+h-ry} V {y+ry} A {rx} {ry} 0 0 1 {x+rx} {y} Z')


def read_vectors(source, color='#173D6A'):
    """Return validated paths, viewBox and provenance. Text is intentionally rejected."""
    source = Path(source)
    raw = source.read_bytes()
    if len(raw) > MAX_SVG_BYTES:
        raise SVGProfileError('SVG exceeds the 2 MB component limit')
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeError as exc:
        raise SVGProfileError('SVG must be UTF-8') from exc
    if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
        raise SVGProfileError('DTD and entity declarations are prohibited')
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise SVGProfileError(f'Malformed SVG: {exc}') from exc
    if root.tag not in ('svg', f'{{{SVG_NS}}}svg'):
        raise SVGProfileError('Expected SVG root')
    elements = list(root.iter())
    if len(elements) > MAX_SVG_ELEMENTS:
        raise SVGProfileError('SVG exceeds the 2000 element component limit')
    viewbox = numbers(root.get('viewBox', '')) or [0, 0, number(root.get('width')), number(root.get('height'))]
    if len(viewbox) != 4 or min(viewbox[2:]) <= 0:
        raise SVGProfileError('A positive width/height or viewBox is required')
    specs, ids = [], set()

    def visit(elem, parent_style, matrix):
        tag = elem.tag.removeprefix(f'{{{SVG_NS}}}')
        if tag in ('title', 'desc'):
            return
        label = elem.get('id', f'part-{len(specs)+1:04}')
        if tag not in ATTRS or (tag == 'svg' and elem is not root):
            raise SVGProfileError(f'{label}: unsupported element {tag}')
        allowed = ATTRS[tag] | STYLE | {'id', 'transform', 'style'}
        unknown = set(elem.attrib) - allowed
        if unknown:
            raise SVGProfileError(f'{label}: unsupported attributes {sorted(unknown)}')
        if elem.get('id'):
            if label in ids:
                raise SVGProfileError(f'Duplicate SVG id: {label}')
            ids.add(label)
        style = dict(parent_style)
        style.update({k: v for k, v in elem.attrib.items() if k in STYLE})
        for field in elem.get('style', '').split(';'):
            if not field.strip():
                continue
            key, sep, value = field.partition(':')
            if not sep or key.strip() not in STYLE:
                raise SVGProfileError(f'{label}: unsupported inline style')
            style[key.strip()] = value.strip()
        if style.get('fill-rule', 'nonzero') != 'nonzero':
            raise SVGProfileError(f'{label}: even-odd fill must be normalized first')
        matrix = multiply(matrix, transform(elem.get('transform', '')))
        if tag in ('svg', 'g'):
            for child in elem:
                visit(child, style, matrix)
            return
        if len(elem):
            raise SVGProfileError(f'{label}: nested content on a primitive is unsupported')
        d = path_for(tag, elem)
        if not d.strip() or len(d) > MAX_PATH_CHARS:
            raise SVGProfileError(f'{label}: empty or excessively long path')
        try:
            commands = normalize_path_commands(svg_path_to_absolute(parse_svg_path(d)))
        except ValueError as exc:
            raise SVGProfileError(f'{label}: {exc}') from exc
        a, b, c, d_, e, f = matrix
        transformed = []
        for command in commands:
            pts = []
            for x, y in zip(command.args[::2], command.args[1::2]):
                pts.extend((a*x+c*y+e, b*x+d_*y+f))
            if any(not math.isfinite(v) or abs(v) > 1e9 for v in pts):
                raise SVGProfileError(f'{label}: transformed path exceeds coordinate limits')
            transformed.append(PathCommand(command.cmd, pts))
        cap, join = style.get('stroke-linecap', 'butt'), style.get('stroke-linejoin', 'miter')
        if cap not in ('butt', 'round', 'square') or join not in ('miter', 'round', 'bevel'):
            raise SVGProfileError(f'{label}: unsupported line cap/join')
        width = number(style.get('stroke-width'), 1) * math.hypot(a, b)
        if width < 0:
            raise SVGProfileError(f'{label}: negative stroke width')
        alpha = [number(style.get(k), 1) for k in ('fill-opacity', 'stroke-opacity')]
        if any(not 0 <= value <= 1 for value in alpha):
            raise SVGProfileError(f'{label}: opacity outside 0..1')
        current = style.get('color', color)
        fill = paint(style.get('fill', 'black'), current)
        stroke = paint(style.get('stroke', 'none'), current)
        if tag in ('line', 'polyline'):
            if tag == 'line':
                fill = None
        specs.append(dict(id=label, commands=transformed, fill=fill, stroke=stroke,
                          stroke_width=width, cap=cap, join=join, alpha=alpha))

    visit(root, {}, IDENTITY)
    if not specs:
        raise SVGProfileError('SVG contains no supported visible primitives')
    return specs, viewbox, hashlib.sha256(raw).hexdigest()


def path_shape_xml(spec, shape_id, name, ox, oy, scale):
    """Shared native path emitter for SVG solids and manifest component Paint."""
    inner, left, top, w, h = path_commands_to_drawingml(spec['commands'], ox, oy, scale, scale)
    cap = {'butt': 'flat', 'round': 'rnd', 'square': 'sq'}[spec['cap']]
    join = {'miter': '<a:miter lim="400000"/>', 'round': '<a:round/>', 'bevel': '<a:bevel/>'}[spec['join']]
    stroke = spec['stroke'] if spec['stroke_width'] else None
    elem = parse_xml(
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr>'
        f'<p:cNvPr id="{shape_id}" name={quoteattr(name)}/><p:cNvSpPr/><p:nvPr/>'
        f'</p:nvSpPr><p:spPr><a:xfrm><a:off x="{round(left*9525)}" y="{round(top*9525)}"/>'
        f'<a:ext cx="{round(w*9525)}" cy="{round(h*9525)}"/></a:xfrm>'
        '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
        '<a:rect l="l" t="t" r="r" b="b"/><a:pathLst>'
        f'<a:path w="{round(w*9525)}" h="{round(h*9525)}">{inner}</a:path>'
        f'</a:pathLst></a:custGeom>{paint_xml(spec["fill"], spec["alpha"][0])}'
        f'<a:ln w="{round(spec["stroke_width"]*scale*9525)}" cap="{cap}">'
        f'{paint_xml(stroke, spec["alpha"][1])}{join}</a:ln></p:spPr></p:sp>'
    )
    return elem, [left / 96, top / 96, w / 96, h / 96]


def add_svg_component(slide, source, x, y, width, height, *, name='vector', color='#173D6A'):
    """Import a whole validated component as native paths in one editable group.

    Placement uses inches and aspect-preserving contain. SVG text, images, CSS,
    masks, filters, gradients, references and marker elements are not supported.
    """
    if not all(math.isfinite(v) for v in (x, y, width, height)) or min(width, height) <= 0:
        raise SVGProfileError('Placement must be finite with positive dimensions')
    specs, (vx, vy, vw, vh), digest = read_vectors(source, color)
    scale = min(width*96/vw, height*96/vh)
    ox = x*96 + (width*96-vw*scale)/2 - vx*scale
    oy = y*96 + (height*96-vh*scale)/2 - vy*scale
    xml_parts = []
    element_map = []
    next_id = max(int(n.get('id')) for n in slide._element.iter(f'{{{P_NS}}}cNvPr')) + 2
    for index, spec in enumerate(specs):
        elem, bounds = path_shape_xml(spec, next_id + index, name + '/' + spec['id'], ox, oy, scale)
        element_map.append({'source_id': spec['id'], 'shape_id': next_id + index,
                            'output_name': name + '/' + spec['id'],
                            'bounds_inches': bounds})
        xml_parts.append(elem)
    group = slide.shapes.add_group_shape()
    group.name = name
    for elem in xml_parts:
        group.shapes._spTree.append(elem)
    group.shapes._recalculate_extents()
    return {'group': group, 'source_sha256': digest, 'paths': len(specs),
            'source_ids': [spec['id'] for spec in specs], 'raster_count': 0,
            'element_map': element_map, 'group_id': group.shape_id,
            'slide_part': str(slide.part.partname)}


def add_rect_link(slide, source, target, *, source_port='bottom', target_port='top',
                  name='flow-link', color='173D6A', width_pt=1.8):
    """Bind an aligned straight connector to verified rectangular node sites.

    Non-aligned nodes fail before mutation. Use the existing explicit orthogonal
    route contract for multi-bend links, not fabricated universal connection sites.
    """
    if source_port not in PORTS or target_port not in PORTS:
        raise ValueError('Unknown rectangular connection port')
    if source is target:
        raise ValueError('Self-links require an explicit orthogonal route')
    for shape in (source, target):
        if shape._element.getparent() is not slide.shapes._spTree:
            raise ValueError('Bound nodes must be top-level shapes on the destination slide')
        if shape.auto_shape_type not in (MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE) or shape.rotation:
            raise ValueError('Only unrotated rectangular node connection sites are verified')
    def port(shape, side):
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
        return {'top': (x+w//2, y), 'bottom': (x+w//2, y+h),
                'left': (x, y+h//2), 'right': (x+w, y+h//2)}[side]
    start, end = port(source, source_port), port(target, target_port)
    if start == end or (start[0] != end[0] and start[1] != end[1]):
        raise ValueError('Align node ports before creating an axis-aligned bound link')
    # Reject routes that initially run into either attached node.
    direction = (end[0]-start[0], end[1]-start[1])
    outward = {'top': (0, -1), 'bottom': (0, 1), 'left': (-1, 0), 'right': (1, 0)}
    dot = lambda u, v: u[0]*v[0]+u[1]*v[1]
    if dot(direction, outward[source_port]) <= 0 or dot(direction, outward[target_port]) >= 0:
        raise ValueError('Selected ports would route the connector through a node')
    rgb = RGBColor.from_string(color.removeprefix('#'))
    if not math.isfinite(width_pt) or width_pt <= 0:
        raise ValueError('Line width must be positive')
    link = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *start, *end)
    link.name = name
    link.begin_connect(source, PORTS[source_port])
    link.end_connect(target, PORTS[target_port])
    link.line.color.rgb = rgb; link.line.width = Pt(width_pt)
    arrow = OxmlElement('a:tailEnd'); arrow.set('type', 'triangle')
    link._element.spPr.get_or_add_ln().append(arrow)
    return link


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('svg', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--template', type=Path, help='Append one new slide to a COPY of an existing PPTX')
    parser.add_argument('--color', default='#173D6A')
    args = parser.parse_args()
    if args.output.exists() or args.output.suffix.lower() != '.pptx':
        parser.error('Use a new .pptx output path; existing files are never overwritten')
    if args.template and args.template.suffix.lower() != '.pptx':
        parser.error('Only ordinary .pptx templates are supported')
    try:
        specs, viewbox, digest = read_vectors(args.svg, args.color)
        deck = Presentation(str(args.template)) if args.template else Presentation()
        if not args.template:
            deck.slide_width = Inches(8)
            deck.slide_height = Inches(max(1, min(50, 8*viewbox[3]/viewbox[2])))
        noncontent = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
        blank = next((layout for layout in deck.slide_layouts
                      if all(p.placeholder_format.type in noncontent for p in layout.placeholders)), None)
        if blank is None:
            raise SVGProfileError('Template has no blank layout; use the Python component API in your builder')
        slide = deck.slides.add_slide(blank)
        w, h = deck.slide_width/EMU, deck.slide_height/EMU
        if min(w, h) <= 0.8:
            raise SVGProfileError('Slide dimensions are too small for the component margins')
        add_svg_component(slide, args.svg, .4, .4, w-.8, h-.8, color=args.color)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        # Exclusive creation keeps pre-existing outputs safe even under a race.
        with args.output.open('xb') as stream:
            deck.save(stream)
        print(json.dumps({'output': str(args.output), 'slides': len(deck.slides),
                          'paths_added': len(specs), 'source_sha256': digest,
                          'raster_count': 0, 'scope': 'vector components only'}, ensure_ascii=False))
    except (SVGProfileError, OSError, ValueError) as exc:
        parser.exit(2, f'Native vector import failed: {exc}\n')


if __name__ == '__main__':
    main()
