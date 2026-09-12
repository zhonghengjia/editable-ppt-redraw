"""Source-first prepare/build over the existing inventory, extractors and emitters.

No image recognition, renderer, installer, arbitrary callback or old-deck input.
All outputs are new candidates. Human source/actual-render review remains separate.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys

from reconstruction_scene import (SourceFrame, SceneNode, ReconstructionScene,
    assemble_pptx, rectangle_draw, svg_draw, text_draw, manifest_component_draw)

SCRIPTS = Path(__file__).resolve().parent


def load(name):
    spec = importlib.util.spec_from_file_location(name.replace('-', '_'), SCRIPTS / (name + '.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write_json(path, value):
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def _keys(value, required, optional=()):
    if not isinstance(value, dict) or not set(required) <= set(value) or set(value) - set(required) - set(optional):
        raise ValueError('missing or unexpected construction fields')


def _positive(value):
    return type(value) in (int, float) and math.isfinite(value) and value > 0


def _hash(value):
    return isinstance(value, str) and len(value) == 64 and all(c in '0123456789abcdef' for c in value)


def _text(value):
    return isinstance(value, str) and bool(value.strip())


def validate_contract(manifest):
    """Only executable fields; base and feature schemas retain their own validators."""
    if 'construction' not in manifest:
        return []
    try:
        config = manifest['construction']
        _keys(config, ('version', 'slide_width_inches'))
        if type(config['version']) is not int or config['version'] not in (1,2) or not _positive(config['slide_width_inches']):
            raise ValueError('construction version 1/2 and positive slide width required')
        if manifest.get('mode') != 'faithful' or manifest.get('editing_policy', 'native') not in ('native', None):
            raise ValueError('this adapter supports faithful native output only')
        source, canvas = manifest['source'], manifest['canvas']
        if not _hash(source.get('sha256')):
            raise ValueError('source.sha256 required before construction')
        if any(type(source.get(k)) is not int or source[k] <= 0 or canvas.get(k) != source[k]
               for k in ('width', 'height')):
            raise ValueError('adapter requires source-pixel canvas; do not silently rescale observations')
        targets = manifest['targets']
        if len(targets) != 1 or targets[0]['format'] != 'pptx':
            raise ValueError('adapter requires one PPTX target')
        if Path(targets[0]['path']).suffix.lower() != '.pptx':
            raise ValueError('PPTX target filename must end with .pptx')
        items = manifest.get('source_inventory', [])
        if not items:
            raise ValueError('source inventory required before preparation')
        frame = SourceFrame((0, 0, source['width'], source['height']),
                            (0, 0, config['slide_width_inches'], config['slide_width_inches']*source['height']/source['width']))
        by = {i['id']: i for i in items}
        if len(by) != len(items):
            raise ValueError('duplicate source inventory IDs')
        nodes = []
        used_connections = set()
        connections = {c['id']: c for c in manifest.get('connections', [])}
        for item in items:
            identity = item['id']
            recipe = item.get('construction')
            _keys(recipe, ('kind', 'observation'), ('options', 'extract', 'trace', 'asset',
                'parent', 'behind', 'z', 'container', 'connection', 'endpoints'))
            if not _text(recipe['observation']):
                raise ValueError(identity + ': source observation required')
            if item.get('output_name', identity) != identity or item.get('output_slide', 1) != 1 or item.get('required_count', 1) != 1:
                raise ValueError(identity + ': one named editing unit on slide 1 required')
            if item.get('representation') not in ('native_primitive', 'native_composite'):
                raise ValueError(identity + ': adapter does not embed pictures')
            kind = recipe['kind']
            specific = {'text': {'options'}, 'rectangle': {'options'},
                        'svg': {'asset'}, 'trace': {'extract', 'trace'}, 'native_component': set()}
            if kind not in specific:
                raise ValueError(identity + ': unsupported route; no template fallback')
            extra = set(recipe) & {'options', 'extract', 'trace', 'asset'}
            if extra != specific[kind]:
                raise ValueError(identity + ': route fields do not match kind')
            role = 'relation' if item['role'] == 'connector' else ('text' if item['role'] == 'text' else 'component')
            if (role == 'text') != (kind == 'text'):
                raise ValueError(identity + ': text must use native text route')
            if recipe.get('container') is not None:
                if recipe['container'] is not True or role != 'component':
                    raise ValueError(identity + ': invalid container role')
                role = 'container'
            if kind == 'text':
                _keys(recipe['options'], ('font_size_px',), ('color', 'font', 'bold', 'align'))
                text_draw(item['text'], **recipe['options'])
            elif kind == 'rectangle':
                _keys(recipe['options'], (), ('fill', 'stroke', 'stroke_pt'))
                for key in ('fill', 'stroke'):
                    color = recipe['options'].get(key)
                    if color is not None and (not isinstance(color, str) or len(color.lstrip('#')) != 6
                            or any(c not in '0123456789abcdefABCDEF' for c in color.lstrip('#'))):
                        raise ValueError(identity + ': RGB color must contain six hex digits')
                if not _positive(recipe['options'].get('stroke_pt', 1)):
                    raise ValueError(identity + ': positive stroke width required')
                if item['role'] not in ('shape', 'decoration'):
                    raise ValueError(identity + ': rectangle cannot stand in for an icon/chart')
            elif kind == 'svg':
                _keys(recipe['asset'], ('path', 'sha256'))
                if not _text(recipe['asset']['path']) or not _hash(recipe['asset']['sha256']):
                    raise ValueError(identity + ': SVG path/hash required')
            elif kind == 'native_component':
                matches = [c for c in manifest.get('native_components', []) if c['id'] == identity]
                if len(matches) != 1 or matches[0]['output_name'] != identity:
                    raise ValueError(identity + ': native_component requires the same canonical component ID/output_name')
            else:
                extraction = recipe['extract']
                _keys(extraction, ('method', 'foreground', 'background', 'reserved_regions'),
                      ('selector', 'annotations','matting','occlusions'))
                if extraction['method'] not in ('selector', 'grabcut','closed_form') or (extraction['method'] == 'selector') != ('selector' in extraction):
                    raise ValueError(identity + ': explicit extraction method required')
                layered=extraction['method']=='closed_form'
                if layered:
                    if config['version']!=2: raise ValueError('closed_form layers require construction version 2')
                    from source_layers import validate_layer_options
                    validate_layer_options(extraction)
                    if recipe['trace']['representation']=='palette_stack':
                        raise ValueError('soft layers require disjoint palette_edges or smooth; never flatten alpha')
                elif 'matting' in extraction or 'occlusions' in extraction:
                    raise ValueError('layer fields require closed_form extraction')
                if 'annotations' in extraction and extraction['method'] not in ('grabcut','closed_form'):
                    raise ValueError(identity + ': annotations require grabcut/closed_form')
                trace = recipe['trace']
                _keys(trace, ('colors', 'representation', 'max_native_paths', 'max_native_commands'))
                if trace['representation'] not in ('palette_edges', 'palette_stack', 'smooth'):
                    raise ValueError(identity + ': trace must retain source paint')
                if type(trace['colors']) is not int or not 2 <= trace['colors'] <= 64:
                    raise ValueError(identity + ': colors must be 2..64')
                for key in ('max_native_paths', 'max_native_commands'):
                    if type(trace[key]) is not int or not 1 <= trace[key] <= 1_000_000:
                        raise ValueError(identity + ': invalid native edit budget')
                if any(type(v) is not int for v in item['bbox']):
                    raise ValueError(identity + ': extraction context uses integer source pixels')
            endpoints = recipe.get('endpoints')
            if role == 'relation':
                if kind != 'svg' or recipe.get('connection') not in connections:
                    raise ValueError(identity + ': relation requires observed SVG and known connection')
                if not isinstance(endpoints, list) or len(endpoints) != 2 or any(v not in by for v in endpoints):
                    raise ValueError(identity + ': source item endpoints required')
                connection = connections[recipe['connection']]
                if [by[v]['module'] for v in endpoints] != [connection['source'], connection['target']]:
                    raise ValueError(identity + ': item endpoints disagree with module connection')
                if recipe['connection'] in used_connections:
                    raise ValueError('connection emitted more than once')
                used_connections.add(recipe['connection'])
            elif 'connection' in recipe or endpoints is not None:
                raise ValueError(identity + ': endpoints require connector role')
            nodes.append(SceneNode(identity, item['bbox'], lambda *a: None, role,
                recipe.get('parent'), recipe.get('behind', ()), recipe.get('z', 0), endpoints))
        if used_connections != set(connections):
            raise ValueError('all declared connections require source-observed construction')
        if {c['id'] for c in manifest.get('native_components', [])} != {
                i['id'] for i in items if i['construction']['kind'] == 'native_component'}:
            raise ValueError('declared native components must use their canonical construction route')
        ReconstructionScene(frame, nodes).ordered()
        return []
    except (ValueError, TypeError, KeyError, AttributeError) as exc:
        return ['construction: ' + str(exc)]


def runtime_check(manifest, *, node=None, extract_python=None, check_extraction=True):
    required = ('pptx', 'PIL', 'numpy')
    missing = [name for name in required if importlib.util.find_spec(name) is None]
    versions = {}
    for name in required:
        if name not in missing:
            try:
                module = importlib.import_module(name)
                versions[name] = getattr(module, '__version__', 'unknown')
            except ImportError as exc:
                missing.append(name + ': ' + str(exc))
    traces = [i for i in manifest['source_inventory'] if i['construction']['kind'] == 'trace']
    node = node or shutil.which('node')
    if traces and not node:
        missing.append('Node.js (supply --node)')
    if traces and node:
        subprocess.run([str(node), '--version'], check=True, capture_output=True, timeout=15)
    methods={i['construction']['extract']['method'] for i in traces}
    if check_extraction and methods & {'grabcut','closed_form'}:
        interpreter = extract_python or sys.executable
        imports={'PIL','numpy'}
        if 'grabcut' in methods: imports.add('cv2')
        if 'closed_form' in methods: imports.add('scipy')
        check = subprocess.run([str(interpreter), '-c', 'import '+', '.join(sorted(imports))],
            capture_output=True, text=True, timeout=30)
        if check.returncode:
            missing.append('/'.join(sorted(imports))+' in extraction interpreter; use --extract-python')
    if missing:
        raise RuntimeError('Missing installed dependencies: ' + ', '.join(missing) + '; no automatic install/fallback')
    return {'python': sys.executable, 'node': str(node) if node else None,
            'extract_python': str(extract_python or sys.executable), 'versions': versions,
            'pipeline_sha256': sha(Path(__file__)), 'skill_sha256': sha(SCRIPTS.parent / 'SKILL.md')}


def checked_manifest(path):
    data = read_json(path)
    report = load('validate-visual-manifest').validate_manifest(data)
    if not report['valid']:
        raise ValueError('Invalid manifest before construction: ' + '; '.join(report['errors'][:8]))
    if 'construction' not in data:
        raise ValueError('manifest has no executable construction recipes')
    return data


def _bound_file(base, entry):
    path = (base / entry['path']).resolve()
    if sha(path) != entry['sha256']:
        raise ValueError('prepared input hash changed: ' + entry['path'])
    return path


def prepare(manifest_path, output, *, node=None, extract_python=None):
    """Freeze source inventory; extract measured supports and copy observed vectors."""
    from PIL import Image
    manifest_path, output = Path(manifest_path).resolve(), Path(output).resolve()
    if output.exists():
        raise ValueError('use a new preparation directory')
    data = checked_manifest(manifest_path)
    runtime = runtime_check(data, node=node, extract_python=extract_python)
    source = (manifest_path.parent / data['source']['path']).resolve()
    if sha(source) != data['source']['sha256']:
        raise ValueError('original source hash changed')
    with Image.open(source) as image:
        if image.size != (data['source']['width'], data['source']['height']) or getattr(image, 'n_frames', 1) != 1:
            raise ValueError('source dimensions/frame mismatch')
    output.mkdir(parents=True)
    frozen_source = output / ('source' + source.suffix)
    shutil.copyfile(source, frozen_source)
    write_json(output / 'input-manifest.json', data)
    records = []
    try:
        for index, item in enumerate(data['source_inventory']):
            recipe, record = item['construction'], {'id': item['id'], 'kind': item['construction']['kind']}
            folder = output / f'item-{index:04d}'
            if recipe['kind'] == 'trace':
                plan = dict(recipe['extract'], source_sha256=data['source']['sha256'],
                            context_bbox=item['bbox'], observation=recipe['observation'])
                if extract_python:
                    plan_path = output / f'extraction-{index:04d}.json'
                    write_json(plan_path, plan)
                    result = subprocess.run([str(extract_python), '-X', 'utf8', '-B',
                        str(SCRIPTS / 'source_objects.py'), str(frozen_source), str(plan_path), str(folder)],
                        capture_output=True, text=True, encoding='utf-8', timeout=180)
                    if result.returncode:
                        raise RuntimeError(item['id'] + ': extraction failed: ' + result.stderr[-2000:])
                else:
                    from source_objects import extract_source_object
                    outputs, evidence = extract_source_object(frozen_source, plan)
                    folder.mkdir()
                    for name, content in outputs.items():
                        (folder / (name + '.png')).write_bytes(content)
                    write_json(folder / 'candidate.json', evidence)
                evidence = read_json(folder / 'candidate.json')
                # Show the candidate in its original context coordinate system,
                # not a tightly fitted thumbnail that hides crop/placement errors.
                with Image.open(frozen_source) as original:
                    x, y, w, h = item['bbox']
                    context = original.convert('RGBA').crop((x, y, x+w, y+h))
                context.save(folder / 'source-context.png')
                candidate = Image.new('RGBA', (w, h), '#eeeeee')
                bx, by, _, _ = evidence['source_bbox']
                with Image.open(folder / 'rgba.png') as rgba:
                    candidate.alpha_composite(rgba.convert('RGBA'), (bx-x, by-y))
                comparison = Image.new('RGB', (w*2+8, h), '#888888')
                comparison.paste(context.convert('RGB'), (0, 0))
                comparison.paste(candidate.convert('RGB'), (w+8, 0))
                comparison.save(folder / 'comparison.png')
                record['candidate'] = {'path': str((folder / 'candidate.json').relative_to(output)),
                                       'sha256': sha(folder / 'candidate.json')}
                record['source_bbox'] = evidence['source_bbox']
            elif recipe['kind'] == 'svg':
                from native_vectors import read_vectors
                asset = _bound_file(manifest_path.parent, recipe['asset'])
                read_vectors(asset)
                folder.mkdir()
                shutil.copyfile(asset, folder / 'source.svg')
                record['asset'] = {'path': str((folder / 'source.svg').relative_to(output)), 'sha256': sha(asset)}
            records.append(record)
        receipt = {'status': 'PREPARED', 'input_sha256': sha(output / 'input-manifest.json'),
                   'source': {'path': frozen_source.name, 'sha256': sha(frozen_source)},
                   'runtime': runtime, 'items': records,
                   'source_coverage': 'REQUIRES_SOURCE_REVIEW'}
        write_json(output / 'prepared.json', receipt)
        write_json(output / 'review-template.json', {'prepared_sha256': sha(output / 'prepared.json'),
            'items': [{'id': r['id'], 'candidate_sha256': r['candidate']['sha256'],
                       'accepted': False, 'note': ''} for r in records if 'candidate' in r]})
        return receipt
    except Exception as exc:
        write_json(output / 'failure.json', {'status': 'FAILED', 'phase': 'prepare', 'error': str(exc)})
        raise


def build(prepared, output, *, review=None, node=None):
    """Trace qualified support, assemble new deck, reopen and run existing checks."""
    from pptx import Presentation
    from pptx.util import Inches
    from component_fidelity import trace_component
    prepared, output = Path(prepared).resolve(), Path(output).resolve()
    if output.exists():
        raise ValueError('use a new build directory; never overlay a previous deck')
    receipt = read_json(prepared / 'prepared.json')
    if receipt.get('status') != 'PREPARED' or sha(prepared / 'input-manifest.json') != receipt['input_sha256']:
        raise ValueError('prepared manifest changed or preparation incomplete')
    data = checked_manifest(prepared / 'input-manifest.json')
    source = _bound_file(prepared, receipt['source'])
    if sha(source) != data['source']['sha256']:
        raise ValueError('prepared source disagrees with manifest')
    records = receipt['items']
    if [r['id'] for r in records] != [i['id'] for i in data['source_inventory']]:
        raise ValueError('prepared inventory is incomplete or reordered')
    approvals = {}
    if review:
        review_data = read_json(review)
        if review_data.get('prepared_sha256') != sha(prepared / 'prepared.json'):
            raise ValueError('review belongs to a different preparation')
        entries = review_data.get('items', [])
        approvals = {e['id']: e for e in entries}
        if len(approvals) != len(entries):
            raise ValueError('duplicate component reviews')
    if set(approvals) != {r['id'] for r in records if r['kind'] == 'trace'}:
        raise ValueError('inspect extracted components and supply their exact review; no automatic approval')
    # Check all support/asset identities before drawing anything. Evidence hashes
    # bind files, not reviewer truth; no claim of cryptographic attestation.
    for item, record in zip(data['source_inventory'], records):
        if record['kind'] != item['construction']['kind']:
            raise ValueError('prepared route changed')
        if record['kind'] == 'trace':
            path = _bound_file(prepared, record['candidate'])
            evidence = read_json(path)
            approval = approvals[item['id']]
            if approval.get('candidate_sha256') != sha(path):
                raise ValueError(item['id'] + ': stale component inspection')
            if approval.get('accepted') is False:
                raise ValueError(item['id'] + ': component inspection rejected; revise source observations')
            if approval.get('accepted') is not True or not _text(approval.get('note')):
                raise ValueError(item['id'] + ': missing/stale component inspection')
            plan = dict(item['construction']['extract'], source_sha256=data['source']['sha256'],
                        context_bbox=item['bbox'], observation=item['construction']['observation'])
            if (evidence['source_sha256'] != sha(source) or evidence['source_bbox'] != record['source_bbox']
                    or evidence['plan_sha256'] != hashlib.sha256(json.dumps(plan, sort_keys=True).encode()).hexdigest()):
                raise ValueError(item['id'] + ': candidate source/recipe binding changed')
            for key, expected in evidence['outputs_sha256'].items():
                if sha(path.parent / (key + '.png')) != expected:
                    raise ValueError(item['id'] + ': prepared support bytes changed')
        elif record['kind'] == 'svg':
            asset = _bound_file(prepared, record['asset'])
            if sha(asset) != item['construction']['asset']['sha256']:
                raise ValueError(item['id'] + ': SVG differs from source observation')
    # Extraction was already completed. Its separate interpreter is not required
    # during build, but the trace worker and PPTX dependencies are.
    runtime = runtime_check(data, node=node, extract_python=receipt['runtime']['extract_python'], check_extraction=False)
    output.mkdir(parents=True)
    try:
        artifact = output / Path(data['targets'][0]['path']).name
        width = data['construction']['slide_width_inches']
        frame = SourceFrame((0, 0, data['source']['width'], data['source']['height']),
            (0, 0, width, width*data['source']['height']/data['source']['width']))
        nodes = []
        # The frozen manifest retains its original bytes. Resolve source reads in
        # a runtime copy so sampled Paint cannot accidentally read a sibling run.
        drawing_manifest = copy.deepcopy(data)
        drawing_manifest['source']['path'] = str(source)
        for index, (item, record) in enumerate(zip(data['source_inventory'], records)):
            recipe, bbox = item['construction'], item['bbox']
            kind = recipe['kind']
            if kind == 'trace':
                bbox = record['source_bbox']
                evidence_path = prepared / record['candidate']['path']
                evidence = read_json(evidence_path)
                options = recipe['trace']
                part_plan = dict(part_id=item['id'], observation=recipe['observation'],
                    source_sha256=data['source']['sha256'], source_bbox=bbox,
                    max_native_paths=options['max_native_paths'], max_native_commands=options['max_native_commands'])
                svg = output / f'item-{index:04d}.svg'
                if recipe['extract']['method']=='closed_form':
                    paint_input={'source_layer':dict(path=str(evidence_path),sha256=sha(evidence_path))}
                else:
                    paint_input={'source_support':dict(path=str(evidence_path.parent/'alpha.png'),
                        sha256=evidence['outputs_sha256']['alpha'],observation=recipe['observation'])}
                trace_component(source, svg, bbox, colors=options['colors'], node=runtime['node'],
                    representation=options['representation'], part_plan=part_plan, support_paint='source',**paint_input)
                draw = svg_draw(svg, expected_sha256=sha(svg))
            elif kind == 'svg':
                draw = svg_draw(_bound_file(prepared, record['asset']), expected_sha256=record['asset']['sha256'])
            elif kind == 'text':
                draw = text_draw(item['text'], **recipe['options'])
            elif kind == 'native_component':
                draw = manifest_component_draw(drawing_manifest, item['id'], base_dir=prepared)
            else:
                draw = rectangle_draw(**recipe['options'])
            role = 'relation' if item['role'] == 'connector' else ('text' if kind == 'text' else 'component')
            if recipe.get('container'):
                role = 'container'
            nodes.append(SceneNode(item['id'], bbox, draw, role, recipe.get('parent'),
                recipe.get('behind', ()), recipe.get('z', 0), recipe.get('endpoints')))
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(frame.viewport[2]), Inches(frame.viewport[3])
        scene_receipt = assemble_pptx(deck.slides.add_slide(deck.slide_layouts[6]), ReconstructionScene(frame, nodes))
        stream = io.BytesIO()
        deck.save(stream)
        payload = stream.getvalue()
        reopened = Presentation(io.BytesIO(payload))
        actual = list(reopened.slides[0].shapes)
        if [s.name for s in actual] != scene_receipt['paint_order']:
            raise ValueError('saved shape inventory/order mismatch')
        actual_by = {s.name: s for s in actual}
        resolved = copy.deepcopy(data)
        resolved['source']['path'] = str(source)
        resolved['targets'] = [{'format': 'pptx', 'path': str(artifact)}]
        for item in resolved['source_inventory']:
            item.update(output_name=item['id'], output_id=str(actual_by[item['id']].shape_id), output_slide=1)
            if item['role'] == 'text' and actual_by[item['id']].text != item['text']:
                raise ValueError(item['id'] + ': saved text mismatch')
        with artifact.open('xb') as target:
            target.write(payload)
        write_json(output / 'resolved-manifest.json', resolved)
        scene_receipt.update(artifact_sha256=sha(artifact), artifact=artifact.name,
            prepared_sha256=sha(prepared / 'prepared.json'), runtime=runtime,
            status='CANDIDATE', visual_review='NOT_PERFORMED', editor_review='NOT_PERFORMED')
        quality = load('run-quality-checks').run_checks(artifact, output / 'resolved-manifest.json', node=runtime['node'])
        write_json(output / 'quality-report.json', quality)
        scene_receipt.update(automation_passed=quality['automation_passed'], delivery_ready=None,
            pending_checks=[c['check'] for c in quality['checks'] if c['status'] in ('FAIL', 'NOT_VERIFIED')])
        write_json(output / 'build-receipt.json', scene_receipt)
        return scene_receipt
    except Exception as exc:
        write_json(output / 'failure.json', {'status': 'FAILED', 'phase': 'build', 'error': str(exc)})
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    for command in ('doctor', 'prepare', 'build'):
        p = sub.add_parser(command)
        p.add_argument('input', type=Path)
        if command != 'doctor':
            p.add_argument('output', type=Path)
        p.add_argument('--node')
        if command != 'build':
            p.add_argument('--extract-python')
        else:
            p.add_argument('--review', type=Path)
    args = parser.parse_args()
    try:
        if args.command == 'doctor':
            result = runtime_check(checked_manifest(args.input), node=args.node, extract_python=args.extract_python)
        elif args.command == 'prepare':
            result = prepare(args.input, args.output, node=args.node, extract_python=args.extract_python)
        else:
            result = build(args.input, args.output, review=args.review, node=args.node)
        print(json.dumps({'status': result.get('status', 'RUNTIME_AVAILABLE'),
                          'items': len(result.get('items', [])), 'visual_review': 'NOT_PERFORMED',
                          'automation_passed': result.get('automation_passed'),
                          'pending_checks': result.get('pending_checks', []), 'delivery_ready': None}))
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError) as exc:
        print(json.dumps({'status': 'FAILED', 'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
